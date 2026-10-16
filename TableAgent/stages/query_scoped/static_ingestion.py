from __future__ import annotations

import hashlib
import json
import threading
import time as time_module
from dataclasses import dataclass, replace
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from TableAgent.llm import BaseLLM, LLMResponse
from TableAgent.utils.paths import safe_name

from .prompts import (
    CHUNK_SUMMARY_TEMPLATE,
    SHEET_SUMMARY_TEMPLATE,
    STATIC_PROMPT_VERSION,
    SUMMARY_SYSTEM_PROMPT,
    WORKBOOK_SUMMARY_TEMPLATE,
)


_CACHE_LOCKS: dict[str, threading.RLock] = {}
_CACHE_LOCKS_GUARD = threading.Lock()
_STATIC_INGESTION_VERSION = "query-static-ingestion-v3-effective-bounds"


def _cache_lock(key: str) -> threading.RLock:
    with _CACHE_LOCKS_GUARD:
        return _CACHE_LOCKS.setdefault(key, threading.RLock())


@dataclass(frozen=True)
class StaticRegion:
    region_id: str
    workbook_path: Path
    workbook_hash: str
    sheet_name: str
    start_row: int
    end_row: int
    max_column: int
    cell_range: str
    source_path: Path
    summary_path: Path
    summary: str


@dataclass(frozen=True)
class StaticSheetIndex:
    sheet_name: str
    header_context_path: Path
    summary_path: Path
    summary: str
    regions: tuple[StaticRegion, ...]


@dataclass(frozen=True)
class StaticWorkbookIndex:
    workbook_path: Path
    workbook_hash: str
    cache_dir: Path
    manifest_path: Path
    summary_path: Path
    summary: str
    sheets: tuple[StaticSheetIndex, ...]
    cache_hit: bool
    responses: tuple[LLMResponse, ...] = ()

    @property
    def regions(self) -> tuple[StaticRegion, ...]:
        return tuple(region for sheet in self.sheets for region in sheet.regions)


class StaticWorkbookIngestor:
    """Build and load a reusable A1-addressed text index for workbooks."""

    def __init__(self, settings: Any, llm: BaseLLM):
        self.settings = settings
        self.llm = llm

    def prepare(
        self,
        workbook_path: Path,
        *,
        selected_sheets: Iterable[str] = (),
    ) -> StaticWorkbookIndex:
        workbook_path = workbook_path.expanduser().resolve()
        if not workbook_path.is_file():
            raise FileNotFoundError(f"Workbook does not exist: {workbook_path}")
        workbook_hash = self._file_hash(workbook_path)
        recipe = self._recipe(selected_sheets)
        recipe_hash = hashlib.sha256(
            json.dumps(recipe, sort_keys=True, ensure_ascii=False).encode("utf-8")
        ).hexdigest()[:16]
        workbook_cache_dir = (
            self.settings.query_static_cache_dir
            / f"{safe_name(workbook_path.stem)[:80]}_{workbook_hash[:16]}"
        )
        cache_dir = workbook_cache_dir / recipe_hash
        with _cache_lock(str(cache_dir.resolve())):
            cached = self._load(cache_dir, workbook_path, workbook_hash)
            if cached is not None:
                return cached
            if bool(
                getattr(
                    self.settings,
                    "query_static_reuse_compatible_cache",
                    True,
                )
            ):
                cached = self._load_compatible(
                    workbook_cache_dir,
                    workbook_path,
                    workbook_hash,
                    recipe,
                    exclude=cache_dir,
                )
                if cached is not None:
                    return cached
            return self._build(
                workbook_path,
                workbook_hash,
                cache_dir,
                recipe,
                tuple(str(value) for value in selected_sheets if str(value)),
            )

    def _build(
        self,
        workbook_path: Path,
        workbook_hash: str,
        cache_dir: Path,
        recipe: dict[str, Any],
        selected_sheets: tuple[str, ...],
    ) -> StaticWorkbookIndex:
        cache_dir.mkdir(parents=True, exist_ok=True)
        # Query-scoped ingestion needs displayed/cached values, not formula
        # source. Formula-heavy wide sheets can otherwise make even a short
        # row chunk exceed the model context window.
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        responses: list[LLMResponse] = []
        sheet_entries: list[dict[str, Any]] = []
        try:
            available = list(workbook.sheetnames)
            missing = [name for name in selected_sheets if name not in available]
            if missing:
                raise ValueError(
                    f"Selected sheets not found in {workbook_path.name}: "
                    + ", ".join(missing)
                )
            sheet_names = list(selected_sheets) if selected_sheets else available
            for sheet_name in sheet_names:
                worksheet = workbook[sheet_name]
                reported_max_row = max(int(worksheet.max_row or 1), 1)
                max_column = min(
                    max(int(worksheet.max_column or 1), 1),
                    self.settings.query_static_max_columns,
                )
                max_row = self._effective_last_nonempty_row(
                    worksheet,
                    reported_max_row,
                    max_column,
                )
                sheet_dir = cache_dir / "sheets" / safe_name(sheet_name)
                chunks_dir = sheet_dir / "chunks"
                chunks_dir.mkdir(parents=True, exist_ok=True)

                header_end = min(max_row, self.settings.query_static_header_rows)
                header_context = self._cells_markdown(
                    worksheet, 1, header_end, max_column
                )
                header_path = sheet_dir / "header_context.md"
                header_path.write_text(header_context, encoding="utf-8")

                region_entries: list[dict[str, Any]] = []
                chunk_summaries: list[str] = []
                chunk_rows = self.settings.query_static_chunk_rows
                chunk_step = chunk_rows - self.settings.query_static_chunk_overlap_rows
                previous_summary = ""
                chunk_number = 1
                start_row = 1
                while start_row <= max_row:
                    end_row = min(start_row + chunk_rows - 1, max_row)
                    cell_range = (
                        f"A{start_row}:{get_column_letter(max_column)}{end_row}"
                    )
                    region_id = (
                        f"{workbook_hash[:12]}:{safe_name(sheet_name)}:"
                        f"R{start_row}-R{end_row}"
                    )
                    region_cells = self._cells_markdown(
                        worksheet, start_row, end_row, max_column
                    )
                    source_path = chunks_dir / f"chunk_{chunk_number:04d}_source.md"
                    source_text = (
                        f"# Region {region_id}\n\n"
                        f"- Sheet: {sheet_name}\n- Range: {cell_range}\n\n"
                        f"{region_cells}"
                    )
                    source_path.write_text(source_text, encoding="utf-8")
                    response = self._generate_summary(
                        CHUNK_SUMMARY_TEMPLATE.format(
                            workbook_name=workbook_path.name,
                            sheet_name=sheet_name,
                            region_id=region_id,
                            cell_range=cell_range,
                            header_context=header_context,
                            previous_summary=previous_summary,
                            region_cells=region_cells,
                        ),
                        system_prompt=SUMMARY_SYSTEM_PROMPT,
                        context=(
                            f"chunk summary for {workbook_path.name} "
                            f"sheet {sheet_name!r} range {cell_range}"
                        ),
                    )
                    responses.append(response)
                    summary = response.content.strip() or (
                        f"Region {region_id} covers {sheet_name}!{cell_range}."
                    )
                    summary_path = chunks_dir / f"chunk_{chunk_number:04d}_summary.txt"
                    summary_path.write_text(summary + "\n", encoding="utf-8")
                    chunk_summaries.append(
                        f"[{region_id}] {sheet_name}!{cell_range}\n{summary}"
                    )
                    previous_summary = summary
                    region_entries.append(
                        {
                            "region_id": region_id,
                            "sheet_name": sheet_name,
                            "start_row": start_row,
                            "end_row": end_row,
                            "max_column": max_column,
                            "cell_range": cell_range,
                            "source_path": str(source_path.relative_to(cache_dir)),
                            "summary_path": str(summary_path.relative_to(cache_dir)),
                            "summary": summary,
                        }
                    )
                    if end_row >= max_row:
                        break
                    start_row += chunk_step
                    chunk_number += 1

                combined_chunks = self._bounded_join(chunk_summaries)
                sheet_response = self._generate_summary(
                    SHEET_SUMMARY_TEMPLATE.format(
                        workbook_name=workbook_path.name,
                        sheet_name=sheet_name,
                        region_summaries=combined_chunks,
                    ),
                    system_prompt=SUMMARY_SYSTEM_PROMPT,
                    context=(
                        f"sheet summary for {workbook_path.name} "
                        f"sheet {sheet_name!r}"
                    ),
                )
                responses.append(sheet_response)
                sheet_summary = sheet_response.content.strip() or combined_chunks
                sheet_summary_path = sheet_dir / "summary.txt"
                sheet_summary_path.write_text(sheet_summary + "\n", encoding="utf-8")
                sheet_entries.append(
                    {
                        "sheet_name": sheet_name,
                        "reported_max_row": reported_max_row,
                        "effective_max_row": max_row,
                        "header_context_path": str(
                            header_path.relative_to(cache_dir)
                        ),
                        "summary_path": str(
                            sheet_summary_path.relative_to(cache_dir)
                        ),
                        "summary": sheet_summary,
                        "regions": region_entries,
                    }
                )

            sheet_summaries = self._bounded_join(
                [
                    f"## Sheet: {entry['sheet_name']}\n{entry['summary']}"
                    for entry in sheet_entries
                ]
            )
            workbook_response = self._generate_summary(
                WORKBOOK_SUMMARY_TEMPLATE.format(
                    workbook_name=workbook_path.name,
                    sheet_summaries=sheet_summaries,
                ),
                system_prompt=SUMMARY_SYSTEM_PROMPT,
                context=f"workbook summary for {workbook_path.name}",
            )
            responses.append(workbook_response)
            workbook_summary = workbook_response.content.strip() or sheet_summaries
            summary_path = cache_dir / "summary.txt"
            summary_path.write_text(workbook_summary + "\n", encoding="utf-8")
            manifest = {
                "version": 3,
                "complete": True,
                "workbook_path": str(workbook_path),
                "workbook_hash": workbook_hash,
                "recipe": recipe,
                "summary_path": str(summary_path.relative_to(cache_dir)),
                "sheets": sheet_entries,
                "token_usage": {
                    "prompt": sum(item.prompt_tokens for item in responses),
                    "completion": sum(item.completion_tokens for item in responses),
                },
            }
            manifest_path = cache_dir / "index.json"
            manifest_path.write_text(
                json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
            )
        finally:
            workbook.close()
        loaded = self._load(cache_dir, workbook_path, workbook_hash, cache_hit=False)
        if loaded is None:
            raise RuntimeError(f"Failed to build static workbook index: {cache_dir}")
        return replace(loaded, responses=tuple(responses))

    def _load(
        self,
        cache_dir: Path,
        workbook_path: Path,
        workbook_hash: str,
        *,
        cache_hit: bool = True,
    ) -> StaticWorkbookIndex | None:
        manifest_path = cache_dir / "index.json"
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return None
        if not manifest.get("complete") or manifest.get("workbook_hash") != workbook_hash:
            return None
        try:
            summary_path = cache_dir / manifest["summary_path"]
            workbook_summary = summary_path.read_text(encoding="utf-8").strip()
            if self._is_error_summary(workbook_summary):
                return None
            sheets: list[StaticSheetIndex] = []
            for sheet in manifest["sheets"]:
                header_path = cache_dir / sheet["header_context_path"]
                sheet_summary_path = cache_dir / sheet["summary_path"]
                header_path.read_text(encoding="utf-8")
                sheet_summary = sheet_summary_path.read_text(encoding="utf-8").strip()
                if self._is_error_summary(sheet_summary):
                    return None
                regions: list[StaticRegion] = []
                for region in sheet["regions"]:
                    source_path = cache_dir / region["source_path"]
                    region_summary_path = cache_dir / region["summary_path"]
                    source_path.read_text(encoding="utf-8")
                    region_summary = region_summary_path.read_text(
                        encoding="utf-8"
                    ).strip()
                    if self._is_error_summary(region_summary):
                        return None
                    regions.append(
                        StaticRegion(
                            region_id=str(region["region_id"]),
                            workbook_path=workbook_path,
                            workbook_hash=workbook_hash,
                            sheet_name=str(region["sheet_name"]),
                            start_row=int(region["start_row"]),
                            end_row=int(region["end_row"]),
                            max_column=int(region["max_column"]),
                            cell_range=str(region["cell_range"]),
                            source_path=source_path,
                            summary_path=region_summary_path,
                            summary=region_summary,
                        )
                    )
                sheets.append(
                    StaticSheetIndex(
                        sheet_name=str(sheet["sheet_name"]),
                        header_context_path=header_path,
                        summary_path=sheet_summary_path,
                        summary=sheet_summary,
                        regions=tuple(regions),
                    )
                )
        except (KeyError, OSError, TypeError, ValueError):
            return None
        return StaticWorkbookIndex(
            workbook_path=workbook_path,
            workbook_hash=workbook_hash,
            cache_dir=cache_dir,
            manifest_path=manifest_path,
            summary_path=summary_path,
            summary=workbook_summary,
            sheets=tuple(sheets),
            cache_hit=cache_hit,
        )

    def _load_compatible(
        self,
        workbook_cache_dir: Path,
        workbook_path: Path,
        workbook_hash: str,
        recipe: dict[str, Any],
        *,
        exclude: Path,
    ) -> StaticWorkbookIndex | None:
        try:
            candidates = sorted(
                (
                    path
                    for path in workbook_cache_dir.iterdir()
                    if path.is_dir() and path != exclude
                ),
                key=lambda path: path.stat().st_mtime,
                reverse=True,
            )
        except OSError:
            return None

        ignored_recipe_fields = {"chunk_rows", "chunk_overlap_rows"}
        expected = {
            key: value
            for key, value in recipe.items()
            if key not in ignored_recipe_fields
        }
        for candidate in candidates:
            try:
                manifest = json.loads(
                    (candidate / "index.json").read_text(encoding="utf-8")
                )
                candidate_recipe = manifest.get("recipe")
                if not isinstance(candidate_recipe, dict):
                    continue
                comparable = {
                    key: value
                    for key, value in candidate_recipe.items()
                    if key not in ignored_recipe_fields
                }
                if comparable != expected:
                    continue
            except (OSError, json.JSONDecodeError):
                continue
            cached = self._load(candidate, workbook_path, workbook_hash)
            if cached is not None:
                return cached
        return None

    @staticmethod
    def _is_error_summary(content: str) -> bool:
        return content.lstrip().upper().startswith("ERROR:")

    def _generate_summary(
        self,
        prompt: str,
        *,
        system_prompt: str,
        context: str,
    ) -> LLMResponse:
        max_retries = max(
            0,
            int(getattr(self.settings, "query_static_error_retries", 3)),
        )
        retry_delay = max(
            0.0,
            float(
                getattr(
                    self.settings,
                    "query_static_error_retry_delay_seconds",
                    2.0,
                )
            ),
        )
        last_error = "unknown model error"
        for attempt in range(max_retries + 1):
            try:
                response = self.llm.generate(
                    prompt,
                    system_prompt=system_prompt,
                )
            except Exception as exc:
                last_error = f"{type(exc).__name__}: {exc}"
            else:
                content = str(response.content or "")
                if not self._is_error_summary(content):
                    return response
                last_error = content.strip()

            if attempt < max_retries and retry_delay:
                time_module.sleep(retry_delay)

        total_attempts = max_retries + 1
        raise RuntimeError(
            f"LLM failed while generating {context} after "
            f"{total_attempts} attempts: {last_error}"
        )

    def _recipe(self, selected_sheets: Iterable[str]) -> dict[str, Any]:
        return {
            "ingestion_version": _STATIC_INGESTION_VERSION,
            "prompt_version": STATIC_PROMPT_VERSION,
            "model": str(
                getattr(self.llm, "model_name", self.llm.__class__.__name__)
            ),
            "chunk_rows": self.settings.query_static_chunk_rows,
            "chunk_overlap_rows": self.settings.query_static_chunk_overlap_rows,
            "header_rows": self.settings.query_static_header_rows,
            "max_columns": self.settings.query_static_max_columns,
            "max_summary_chars": self.settings.query_static_max_summary_chars,
            "selected_sheets": sorted(
                str(value) for value in selected_sheets if str(value)
            ),
        }

    def _bounded_join(self, sections: list[str]) -> str:
        text = "\n\n".join(sections)
        limit = self.settings.query_static_max_summary_chars
        if len(text) <= limit:
            return text
        return text[:limit] + "\n...TRUNCATED..."

    @staticmethod
    def _effective_last_nonempty_row(
        worksheet: Any,
        reported_max_row: int,
        max_column: int,
    ) -> int:
        """Ignore trailing rows created only by spreadsheet formatting."""
        last_nonempty_row = 0
        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=reported_max_row,
                min_col=1,
                max_col=max_column,
                values_only=True,
            ),
            1,
        ):
            if any(value is not None for value in values):
                last_nonempty_row = row_number
        return max(last_nonempty_row, 1)

    @staticmethod
    def _cells_markdown(
        worksheet: Any,
        start_row: int,
        end_row: int,
        max_column: int,
    ) -> str:
        lines = ["| Cell | Value |", "|---|---|"]
        for row in worksheet.iter_rows(
            min_row=start_row,
            max_row=end_row,
            min_col=1,
            max_col=max_column,
        ):
            for cell in row:
                if cell.value is None:
                    continue
                value = StaticWorkbookIngestor._format_value(cell.value)
                lines.append(f"| {cell.coordinate} | {value} |")
        if len(lines) == 2:
            lines.append("| (empty) | |")
        return "\n".join(lines)

    @staticmethod
    def _format_value(value: Any) -> str:
        if isinstance(value, (datetime, date, time)):
            text = value.isoformat()
        else:
            text = str(value)
        return (
            text.replace("\\", "\\\\")
            .replace("|", "\\|")
            .replace("\r", " ")
            .replace("\n", "<br>")
        )

    @staticmethod
    def _file_hash(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()
