from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

import openpyxl

from TableAgent.pipeline.sample import EvalSample
from TableAgent.pipeline.component import RuntimeComponent
from TableAgent.pipeline.contracts import PipelineRuntimeContract
from TableAgent.shared.artifacts import prepared_verification
from TableAgent.pipeline.sample import has_workbook_sources

from .cache import StructureCacheRecord
from .contracts import StructureInput
from .metadata import SheetMetadata
from .source_preparer import SourcePreparer
from TableAgent.utils.paths import safe_name


class StructurePipeline(RuntimeComponent):
    """Coordinate structure work through the pipeline runtime contract."""

    def __init__(self, runtime: PipelineRuntimeContract):
        super().__init__(runtime)

    def verify_samples(
        self,
        samples: list[EvalSample],
        *,
        force: bool = True,
    ) -> list[StructureCacheRecord]:
        output = self.stages.structure.run(
            StructureInput(samples=tuple(samples), force=force)
        )
        return list(output.records)

    def _verify_samples_impl(
        self,
        samples: list[EvalSample],
        *,
        force: bool = True,
    ) -> list[StructureCacheRecord]:
        source_samples = [
            sample
            for sample in samples
            if has_workbook_sources(sample) and self.settings.should_retrieve(sample)
        ]
        standard_samples = [
            sample
            for sample in samples
            if not has_workbook_sources(sample) or not self.settings.should_retrieve(sample)
        ]
        records = []
        for sample in standard_samples:
            # StructureCache.prepare already emits structure_done with the real
            # source workbook name. Do not emit again here: record.workbook_path
            # is often the staged copy named workbook.xlsx and double-counts.
            record = self.structure_cache.prepare(sample, force=force)
            records.append(record)
        self._verified_samples.update(
            {
                sample.sample_id: record
                for sample, record in zip(standard_samples, records)
            }
        )
        if source_samples:
            if not self.settings.perfect_retrieval:
                self.source_preparer.prepare(
                    source_samples,
                    regenerate_invalid=force,
                    force=force,
                )
            seen: set[Path] = set()
            for sample in source_samples:
                candidates = (
                    self.source_retriever.load_perfect_candidates(sample)
                    if self.settings.perfect_retrieval
                    else self.source_retriever.load_candidates(sample)
                )
                if not candidates:
                    failure_dir = (
                        self.settings.source_artifact_dir
                        or self.settings.structure_cache_dir
                    )
                    table_part = safe_name(sample.table_id or sample.sample_id)[:80] or "table"
                    digest = hashlib.sha256(sample.sample_id.encode("utf-8")).hexdigest()[:8]
                    key = f"{table_part}_{digest}"
                    records.append(
                        StructureCacheRecord(
                            key=key,
                            directory=failure_dir,
                            workbook_path=Path(str(sample.table_path).split(";")[0]),
                            sheet_name="",
                            structure_path=failure_dir / "structure.yaml",
                            manifest_path=failure_dir / "metadata.json",
                            status="not_good",
                            cache_hit=False,
                        )
                    )
                    continue
                for candidate in candidates:
                    if candidate.directory in seen:
                        continue
                    seen.add(candidate.directory)
                    workbook_stem = (
                        Path(candidate.workbook_path).stem
                        if candidate.workbook_path
                        else Path(candidate.directory).parent.name
                    )
                    sheet_part = safe_name(candidate.sheet_name or candidate.directory.name)[:40] or "sheet"
                    table_part = safe_name(workbook_stem or sample.table_id or "table")[:80] or "table"
                    digest = hashlib.sha256(
                        str(candidate.directory.resolve()).encode("utf-8")
                    ).hexdigest()[:8]
                    key = f"{table_part}_{sheet_part}_{digest}"
                    verification = prepared_verification(candidate.directory)
                    records.append(
                        StructureCacheRecord(
                            key=key,
                            directory=candidate.directory,
                            workbook_path=candidate.workbook_path,
                            sheet_name=candidate.sheet_name,
                            structure_path=candidate.directory / "structure.yaml",
                            manifest_path=candidate.directory / "metadata.json",
                            status=str(verification.get("status") or "good"),
                            cache_hit=not force,
                        )
                    )
        return records

    @staticmethod
    def structure_progress_totals(samples: list[EvalSample]) -> dict[str, Any]:
        """Count the structure work units shown by the CLI progress bar."""
        standard_samples = [
            sample for sample in samples if not has_workbook_sources(sample)
        ]
        source_samples = [sample for sample in samples if has_workbook_sources(sample)]
        selected_sheets = set(SourcePreparer.selected_sheet_names(source_samples))
        sheets_per_file = {
            f"sample:{sample.sample_id}": 1 for sample in standard_samples
        }
        files_per_key = {key: 1 for key in sheets_per_file}

        for source_path in SourcePreparer._source_paths(source_samples):
            try:
                workbook = openpyxl.load_workbook(
                    source_path,
                    read_only=True,
                    data_only=True,
                )
                try:
                    sheet_count = max(
                        len(
                            [
                                name
                                for name in workbook.sheetnames
                                if not selected_sheets or name in selected_sheets
                            ]
                        ),
                        1,
                    )
                finally:
                    workbook.close()
            except Exception:
                sheet_count = 1
            key = f"book:{source_path.name}"
            sheets_per_file[key] = sheets_per_file.get(key, 0) + sheet_count
            files_per_key[key] = files_per_key.get(key, 0) + 1

        return {
            "files": sum(files_per_key.values()),
            "sheets": sum(sheets_per_file.values()),
            "sheets_per_file": sheets_per_file,
            "files_per_key": files_per_key,
        }

    def _analyze_source_sheet(
        self,
        source_path: Path,
        sheet_name: str,
        metadata: SheetMetadata,
        sheet_dir: Path,
    ) -> str:
        if self.layout_workflow is None:
            raise RuntimeError("Source verification requires a layout VLM client")
        self._progress(
            "prepare",
            workbook=source_path.name,
            sheet=sheet_name,
            range=metadata.used_range,
        )
        result = self.layout_workflow.run(
            workbook_path=source_path,
            sheet_name=sheet_name,
            metadata=metadata,
            output_dir=sheet_dir,
        )
        metadata_path = sheet_dir / "metadata.json"
        if metadata_path.is_file():
            metadata_payload = json.loads(metadata_path.read_text(encoding="utf-8"))
            metadata_payload["verification"] = result.verification
            metadata_path.write_text(
                json.dumps(metadata_payload, ensure_ascii=False),
                encoding="utf-8",
            )
        return result.structure_text

    @staticmethod
    def _metadata_for_workbook_sheet(
        workbook_path: Path,
        sheet_name: str,
    ) -> SheetMetadata:
        workbook = openpyxl.load_workbook(
            workbook_path,
            read_only=False,
            data_only=False,
        )
        try:
            worksheet = workbook[sheet_name]
            used_range = worksheet.calculate_dimension()
            merged_ranges = [
                str(cell_range) for cell_range in worksheet.merged_cells.ranges
            ]
            if used_range == "A1:A1" and worksheet["A1"].value is None:
                used_range = None
            return SheetMetadata(sheet_name, used_range, merged_ranges)
        finally:
            workbook.close()
