from copy import copy

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment

from TableAgent.stages.structure.compression.helper import SheetHelper
from TableAgent.stages.structure.compression.formater import SheetFormater


class SheetCompressor:
    def __init__(self, min_row_to_compress: int = 50):
        self.min_row_to_compress = min_row_to_compress

    def _get_cell_style_fp(self, cell):
        is_b = cell.font.b
        is_i = cell.font.i
        is_u = cell.font.u
        text_color = cell.font.color.rgb if cell.font.color is not None else None
        background_color = (
            cell.fill.fgColor.rgb if cell.fill.fgColor is not None else None
        )
        return (
            is_b, is_i, is_u, 
            text_color, 
            background_color
        )

    def _get_cell_data(self, cell):
        return cell.value

    def _get_cell_dtype_fp(self, cell):
        data = self._get_cell_data(cell)
        if data is None or (isinstance(data, str) and not data.strip()):
            return "empty"
        if isinstance(data, (int, float)):
            return "number"
        return "string"

    def _get_cell_fp(self, cell):
        return [self._get_cell_style_fp(cell), self._get_cell_dtype_fp(cell)]

    def _get_row_fp(self, ws, row_idx):
        _, ncols = SheetHelper.get_n_row_col(ws)
        row_fp = []
        for col_idx in range(2, ncols + 1):  # bỏ cột index
            cell = ws.cell(row_idx, col_idx)
            row_fp.append(self._get_cell_fp(cell))
        return row_fp

    def _capture_row_styles(self, ws, row_idx, ncols):
        styles = []
        for col_idx in range(1, ncols + 1):
            cell = ws.cell(row_idx, col_idx)
            styles.append(
                {
                    "font": copy(cell.font),
                    "fill": copy(cell.fill),
                    "border": copy(cell.border),
                    "alignment": copy(cell.alignment),
                    "number_format": cell.number_format,
                    "protection": copy(cell.protection),
                }
            )
        return styles

    def _set_ellipsis_row(self, ws, row_idx, ncols, row_styles):
        for col_idx in range(1, ncols + 1):
            cell = ws.cell(row=row_idx, column=col_idx)

            # Only the top-left cell of a merged range is writable.  The
            # inserted ellipsis row can cross a vertical merge, in which case
            # openpyxl returns a read-only MergedCell placeholder here.  Keep
            # that merge intact and write ellipses only to normal cells.
            if isinstance(cell, MergedCell):
                continue

            style = row_styles[col_idx - 1]

            cell.font = copy(style["font"])
            cell.fill = copy(style["fill"])
            cell.border = copy(style["border"])
            cell.number_format = style["number_format"]
            cell.protection = copy(style["protection"])

            old_align = style["alignment"]
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=old_align.wrap_text,
            )

            cell.value = "..."

    def vertical_compress(self, file_path: str, sheet_name: str, output_path: str = "abc.xlsx", k: int = 2):
        wb = load_workbook(file_path)
        ws = wb[sheet_name]

        SheetFormater.add_sheet_index_cells(ws)
        SheetFormater.fit_col_width(ws)

        nrows, ncols = SheetHelper.get_n_row_col(ws)
        origin_rows = nrows

        if nrows < self.min_row_to_compress:
            print(f"Visual Compress Rows: {origin_rows} -> {origin_rows} (-0.00%)")
            wb.save(output_path)
            return ws

        row_idx = 2

        while row_idx <= nrows:
            row_fp = self._get_row_fp(ws, row_idx)
            end_row_idx = row_idx
            while end_row_idx + 1 <= nrows:
                next_row_fp = self._get_row_fp(ws, end_row_idx + 1)
                if next_row_fp != row_fp: break
                end_row_idx += 1

            n_same_rows = end_row_idx - row_idx + 1
            if n_same_rows > 2 * k:
                delete_start = row_idx + k
                delete_count = n_same_rows - 2 * k

                row_styles = self._capture_row_styles(ws, delete_start, ncols)
                SheetHelper.delete_rows(ws, delete_start, delete_count)
                SheetHelper.insert_rows(ws, delete_start, 1)
                self._set_ellipsis_row(ws, delete_start, ncols, row_styles)

                nrows = nrows - delete_count + 1
                row_idx = delete_start + k + 1

            else:
                row_idx = end_row_idx + 1

        after_rows, _ = SheetHelper.get_n_row_col(ws)
        reduction_percent = ((origin_rows - after_rows) / origin_rows * 100 if origin_rows > 0 else 0)
        print(f"Visual Compress Rows: {origin_rows} -> {after_rows} (-{reduction_percent:.2f}%)")
        wb.save(output_path)
        return ws
