from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import PatternFill, Alignment, Font, Side, Border
from openpyxl.utils import get_column_letter

from TableAgent.stages.structure.compression.helper import SheetHelper

class SheetFormater:
    @staticmethod
    def add_sheet_index_cells(ws: Worksheet):
        merge_ranges = list(ws.merged_cells.ranges)
        
        # unmerge
        for rng in merge_ranges:
            ws.unmerge_cells(str(rng))
        
        # insert index row, col
        ws.insert_rows(1)
        ws.insert_cols(1)
        
        # re-merge
        for rng in merge_ranges:
            ws.merge_cells(
                start_row=rng.min_row + 1,
                end_row=rng.max_row + 1,
                start_column=rng.min_col + 1,
                end_column=rng.max_col + 1
            )
        
        # index cell style
        gray_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        bold_font = Font(bold=True)
        thin_side = Side(style="thin", color="B0B0B0")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        def _style_cell(cell):
            cell.fill = gray_fill
            cell.alignment = center_align
            cell.font = bold_font
            cell.border = thin_border
        
        # add index value
        _style_cell(ws.cell(1, 1))
        
        nrows, ncols = SheetHelper.get_n_row_col(ws)
        for col_idx in range(2, ncols + 2):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = get_column_letter(col_idx - 1)
            _style_cell(cell)
            
        for row_idx in range(2, nrows + 2):
            cell = ws.cell(row=row_idx, column=1)
            cell.value = row_idx - 1
            _style_cell(cell)

    @staticmethod
    def fit_col_width(ws: Worksheet):
        MAX_WIDTH = 50
        MIN_WIDTH = 10
        WIDTH_PADDING = 3
        
        nrows, ncols = SheetHelper.get_n_row_col(ws)
        for col_idx in range(1, ncols + 2):
            max_length = 0
            for row_idx in range(1, nrows + 2):
                cell_value = ws.cell(row_idx, col_idx).value
                if cell_value is not None: max_length = max(max_length, len(str(cell_value)))
        
            width = min(max(max_length + WIDTH_PADDING, MIN_WIDTH), MAX_WIDTH)
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width
        
            if max_length + WIDTH_PADDING > MAX_WIDTH:
                for row_idx in range(1, nrows + 2):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    old_align = cell.alignment
                    cell.alignment = Alignment(
                        horizontal=old_align.horizontal,
                        vertical=old_align.vertical or "top",
                        wrap_text=True,
                    )
    