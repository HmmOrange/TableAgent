from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from TableAgent.stages.structure.structure_prompts import (
    LAYOUT_MAS_SYSTEM_PROMPT,
    LAYOUT_MAS_USER_PROMPT_TEMPLATE,
)
from TableAgent.shared.agents import AgentMessage, BaseTableAgent
from TableAgent.llm import BaseLLM, LLMResponse

from TableAgent.stages.structure.layout.parsing import (
    _is_valid_structure,
    extract_layout_structure_result,
)

@dataclass(frozen=True)
class LayoutResult:
    structure_text: str
    changelog: str
    changed: bool
    response: LLMResponse
    discarded: str
    preflight_errors: list[str]


class LayoutAgent(BaseTableAgent):
    name = "LayoutAgent"
    profile = "Spreadsheet layout VLM"
    goal = "Extract all visible table headers and data ranges from the complete used range."

    def __init__(self, vlm: BaseLLM):
        super().__init__()
        self.vlm = vlm

    def run(
        self,
        *,
        metadata_text: str,
        structure_text: str,
        image_path: Path,
        sheet_range: str,
        feedback: str,
        iteration: int,
        iteration_dir: Path,
    ) -> LayoutResult:
        feedback_block = ""
        if feedback:
            # Lead with a short correction rule so weaker models do not anchor on
            # the rejected coordinate in the previous YAML.
            feedback_block = (
                "\nRETRY THIS CORRECTION:\n"
                "The ranges named below are wrong. Do not use them again. "
                "Find each label in the image and try again. "
                "Change every flagged field before returning YAML.\n\n"
                "Deterministic verifier feedback:\n"
                f"{feedback}\n"
            )
        prompt = LAYOUT_MAS_USER_PROMPT_TEMPLATE.format(
            metadata_text=metadata_text,
            sheet_range=sheet_range,
            structure_text=structure_text or "{}",
            feedback_block=feedback_block,
        )
        iteration_dir.joinpath("layout_prompt.txt").write_text(prompt, encoding="utf-8")
        response = self.vlm.generate_with_image(
            prompt=prompt,
            image_path=image_path,
            system_prompt=LAYOUT_MAS_SYSTEM_PROMPT,
        )
        iteration_dir.joinpath("layout_response.txt").write_text(response.content, encoding="utf-8")
        parsed = extract_layout_structure_result(response.content)
        updated = parsed.structure_text
        group_errors: list[str] = []
        if not _is_valid_structure(updated):
            updated = structure_text
        else:
            updated, group_errors = _merge_existing_structure(structure_text, updated)
        changed = bool(updated.strip()) and _canonical_yaml(updated) != _canonical_yaml(structure_text)
        changelog = parsed.changelog or ("Structure updated." if changed else "No change.")
        if not changed:
            changelog = "No change."
        self.remember(AgentMessage(
            sent_from=self.name,
            sent_to="deterministic_verifier",
            content=changelog,
            iteration=iteration,
            metadata={"sheet_range": sheet_range, "changed": changed},
        ))
        return LayoutResult(
            updated,
            changelog,
            changed,
            response,
            parsed.discarded,
            parsed.preflight_errors + group_errors,
        )


def _merge_existing_structure(previous_text: str, updated_text: str) -> tuple[str, list[str]]:
    merged = _union_existing_data_ranges(previous_text, updated_text)
    try:
        previous = yaml.safe_load(previous_text) if previous_text.strip() else None
        updated = yaml.safe_load(merged) if merged.strip() else None
    except yaml.YAMLError:
        return merged, []
    if not isinstance(previous, dict) or not isinstance(updated, dict):
        return merged, []
    errors: list[str] = []
    changed = _merge_existing_groups(previous, updated, errors)
    if not changed:
        return merged, errors
    return yaml.safe_dump(updated, sort_keys=False, allow_unicode=True).strip(), errors


def _merge_existing_groups(
    previous: dict[str, Any],
    updated: dict[str, Any],
    preflight_errors: list[str] | None = None,
) -> bool:
    changed = False
    for table_id, previous_table in previous.items():
        updated_table = updated.get(table_id)
        if not isinstance(previous_table, dict) or not isinstance(updated_table, dict):
            continue
        old_groups = previous_table.get("groups") or []
        new_groups = updated_table.get("groups") or []
        if not isinstance(old_groups, list) or not isinstance(new_groups, list):
            continue
        new_by_id = {
            str(group.get("id")): group
            for group in new_groups
            if isinstance(group, dict) and group.get("id")
        }
        for old_group in old_groups:
            if not isinstance(old_group, dict) or not old_group.get("id"):
                continue
            group_id = str(old_group["id"])
            new_group = new_by_id.get(group_id)
            if new_group is None:
                new_groups.append(dict(old_group))
                changed = True
                continue
            for field in ("label", "description", "axis"):
                if not new_group.get(field) and old_group.get(field):
                    new_group[field] = old_group[field]
                    changed = True
            for field in ("group_range", "data_range"):
                old_range = old_group.get(field)
                new_range = new_group.get(field)
                if not new_range and old_range:
                    new_group[field] = old_range
                    changed = True
                    continue
                if not old_range or not new_range:
                    continue
                try:
                    old_box = range_boundaries(str(old_range))
                    new_box = range_boundaries(str(new_range))
                except (TypeError, ValueError):
                    continue
                if _boxes_overlap_or_touch(old_box, new_box):
                    unioned = _box_to_range((
                        min(old_box[0], new_box[0]), min(old_box[1], new_box[1]),
                        max(old_box[2], new_box[2]), max(old_box[3], new_box[3]),
                    ))
                    if unioned != new_range:
                        new_group[field] = unioned
                        changed = True
                elif preflight_errors is not None:
                    preflight_errors.append(
                        f"{table_id}.groups[{group_id}].{field}: disjoint viewport ranges; "
                        "use separate group IDs or correct the range"
                    )
        old_ids = [
            str(group.get("id"))
            for group in old_groups
            if isinstance(group, dict) and group.get("id")
        ]
        reordered = []
        for group_id in old_ids:
            match = next(
                (group for group in new_groups if isinstance(group, dict) and str(group.get("id")) == group_id),
                None,
            )
            if match is not None:
                reordered.append(match)
        reordered.extend(
            group
            for group in new_groups
            if not isinstance(group, dict) or str(group.get("id")) not in old_ids
        )
        if reordered != new_groups:
            new_groups[:] = reordered
            changed = True
        updated_table["groups"] = new_groups
    return changed


def _canonical_yaml(text: str) -> Any:
    if not text.strip():
        return None
    try:
        return yaml.safe_load(text)
    except yaml.YAMLError:
        return text.strip()


def _union_existing_data_ranges(previous_text: str, updated_text: str) -> str:
    try:
        previous = yaml.safe_load(previous_text) if previous_text.strip() else None
        updated = yaml.safe_load(updated_text) if updated_text.strip() else None
    except yaml.YAMLError:
        return updated_text
    if not isinstance(previous, dict) or not isinstance(updated, dict):
        return updated_text

    previous_headers = {
        _header_identity(path, header): header
        for path, header in _iter_structure_headers(previous)
    }
    changed = False
    for path, header in _iter_structure_headers(updated):
        prior = previous_headers.get(_header_identity(path, header))
        if not prior:
            continue
        old_range = prior.get("data_range")
        new_range = header.get("data_range")
        orientation = str(header.get("orientation") or prior.get("orientation") or "column").lower()
        unioned = _union_data_range(old_range, new_range, orientation)
        if unioned and unioned != new_range:
            header["data_range"] = unioned
            changed = True
    if not changed:
        return updated_text
    return yaml.safe_dump(updated, sort_keys=False, allow_unicode=True).strip()


def _iter_structure_headers(structure: dict[str, Any]):
    for table_key, table in structure.items():
        if not isinstance(table, dict):
            continue
        headers = table.get("headers") or []
        if not isinstance(headers, list):
            continue
        for index, header in enumerate(headers):
            yield from _iter_header_tree(header, f"{table_key}.headers[{index}]")


def _iter_header_tree(header: Any, path: str):
    if not isinstance(header, dict):
        return
    yield path, header
    sub_headers = header.get("sub_headers") or []
    if not isinstance(sub_headers, list):
        return
    for index, child in enumerate(sub_headers):
        yield from _iter_header_tree(child, f"{path}.sub_headers[{index}]")


def _header_identity(path: str, header: dict[str, Any]) -> tuple[str, str, str]:
    return (
        path,
        str(header.get("label") or "").strip().casefold(),
        str(header.get("header_range") or header.get("range") or "").strip().upper(),
    )


def _union_data_range(old_range: Any, new_range: Any, orientation: str) -> str | None:
    if not old_range or not new_range:
        return None
    try:
        old_box = range_boundaries(str(old_range))
        new_box = range_boundaries(str(new_range))
    except (TypeError, ValueError):
        return None

    if orientation == "row":
        if old_box[1] != new_box[1] or old_box[3] != new_box[3]:
            return None
    else:
        if old_box[0] != new_box[0] or old_box[2] != new_box[2]:
            return None
    if not _boxes_overlap_or_touch(old_box, new_box):
        return None
    return _box_to_range((
        min(old_box[0], new_box[0]),
        min(old_box[1], new_box[1]),
        max(old_box[2], new_box[2]),
        max(old_box[3], new_box[3]),
    ))


def _boxes_overlap_or_touch(left: tuple[int, int, int, int], right: tuple[int, int, int, int]) -> bool:
    return not (
        left[2] + 1 < right[0]
        or right[2] + 1 < left[0]
        or left[3] + 1 < right[1]
        or right[3] + 1 < left[1]
    )


def _box_to_range(box: tuple[int, int, int, int]) -> str:
    min_col, min_row, max_col, max_row = box
    start = f"{get_column_letter(min_col)}{min_row}"
    end = f"{get_column_letter(max_col)}{max_row}"
    return start if start == end else f"{start}:{end}"


__all__ = ["LayoutAgent", "LayoutResult"]

