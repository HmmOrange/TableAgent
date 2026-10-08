from __future__ import annotations

import json
import re
from typing import Any

from TableAgent.stages.qa.actions.base_action import ReviewResult
from TableAgent.stages.qa.actions.review import _salvage_review
from TableAgent.stages.qa.actions.understand_question import _cell_text, decisions_block
from TableAgent.stages.qa.prompts.understanding import DECISIONS_FOR_FINAL_REVIEW
from TableAgent.stages.qa.prompts.review import (
    FINAL_ANSWER_REVIEW_SYSTEM_PROMPT,
    FINAL_ANSWER_REVIEW_USER_PROMPT_TEMPLATE,
)


A1_REFERENCE = re.compile(
    r"(?:'(?P<quoted>[^']+)'!|(?P<sheet>[A-Za-z_][\w.]*)!)?\b(?P<col>[A-Z]{1,3})(?P<row>[1-9]\d{0,5})\b"
)
CELL_REPR = re.compile(r"Cell\(row=(?P<row>\d+), col=(?P<col>\d+)\)")
PERCENT_DECIMALS = re.compile(r"\.([0#]+)%")
MAX_CONTEXT_ANCHORS = 6
MAX_CONTEXT_CHARS = 6000
CONTEXT_ROWS = 2  # rows shown above and below each cell the code used
CONTEXT_COLUMNS = 1  # columns shown left and right of it


class ReviewFinalAnswerAction:
    """Independently verify the final answer against successful runtime evidence."""

    def __init__(self, env: Any, llm_client: Any | None = None):
        self.env = env
        self.llm_client = llm_client

    def run(self, *, question: str, plan: list[Any], outputs: list[Any], final_answer: str) -> ReviewResult:
        if self.llm_client is None:
            return ReviewResult(accepted=True, feedback="Final review skipped without an LLM client.", score=1.0)

        plan_text = json.dumps([
            {
                "id": subtask.id,
                "layer": subtask.layer,
                "description": subtask.description,
                "depends_on": list(subtask.depends_on),
            }
            for subtask in plan
        ], ensure_ascii=False, indent=2)
        evidence_sections = []
        for output in outputs:
            if not output.success:
                continue
            observation = str(output.observation or "")
            code = str(output.code or "")
            evidence_sections.append(
                f"## {output.subtask_id} [{output.layer}]\n"
                f"Code:\n```python\n{code[-3000:]}\n```\n"
                f"Observation:\n{observation[:4000]}"
            )
        evidence = "\n\n".join(evidence_sections) or "No successful runtime evidence was produced."
        cell_context = self._cell_context(outputs)
        grouped_headers = self._grouped_header_context()
        groups = self._group_context(question)
        prompt = FINAL_ANSWER_REVIEW_USER_PROMPT_TEMPLATE.format(
            question=question,
            question_decisions=decisions_block(self.env, DECISIONS_FOR_FINAL_REVIEW),
            plan=plan_text,
            evidence=evidence,
            cell_context=cell_context,
            grouped_headers=grouped_headers,
            groups=groups,
            final_answer=final_answer,
        )
        self.env.logger.log_event("final_answer_review_prompt", {
            "prompt": prompt,
            "system_prompt": FINAL_ANSWER_REVIEW_SYSTEM_PROMPT,
        })
        try:
            response = self.llm_client.generate(prompt, system_prompt=FINAL_ANSWER_REVIEW_SYSTEM_PROMPT)
        except Exception as exc:
            self.env.logger.log_event("final_answer_review_error", {"error": str(exc)})
            return ReviewResult(accepted=True, feedback=f"Final review unavailable: {exc}", score=0.0)

        self.env.logger.log_event("final_answer_review_response", {"content": response.content})
        json_match = re.search(r"```json\s*(.*?)\s*```", str(response.content), re.DOTALL)
        payload = json_match.group(1) if json_match else str(response.content).strip()
        try:
            data = json.loads(payload)
        except (TypeError, ValueError, json.JSONDecodeError):
            data = _salvage_review(str(response.content))
            if data is None:
                return ReviewResult(accepted=True, feedback="Final reviewer returned invalid JSON; review skipped.", score=0.0)
        if not isinstance(data, dict):
            return ReviewResult(accepted=True, feedback="Final reviewer returned a non-object; review skipped.", score=0.0)
        if "accepted" not in data:
            return ReviewResult(accepted=True, feedback="Final reviewer omitted `accepted`; review skipped.", score=0.0)
        accepted_value = data.get("accepted")
        if not isinstance(accepted_value, bool):
            return ReviewResult(
                accepted=True,
                feedback="Final reviewer returned a non-boolean `accepted`; review skipped.",
                score=0.0,
            )
        accepted = accepted_value
        try:
            score = max(0.0, min(1.0, float(data.get("score", 1.0 if accepted else 0.0))))
        except (TypeError, ValueError):
            score = 1.0 if accepted else 0.0
        feedback = str(data.get("feedback") or ("Accepted." if accepted else "Rejected.")).strip()
        cause = str(data.get("cause") or "").strip().lower()
        cause = "interpretation" if cause == "interpretation" and not accepted else "execution"
        return ReviewResult(accepted=accepted, feedback=feedback, score=score, cause=cause)

    def _cell_context(self, outputs: list[Any]) -> str:
        """The worksheet around the cells the accepted code referenced.

        Code shows which cells it read but not their surroundings: the column headers above,
        the section label a row belongs to, or whether a neighboring row is a total. Showing a
        small window lets the reviewer check those against the question.
        """
        workbook = getattr(self.env, "workbook", None)
        if workbook is None:
            return "No worksheet context was available."
        candidates: list[tuple[str, int, int]] = []
        accepted = [output for output in outputs if output.success]
        # Addresses written in code come first; printed ones cover code that read through operators.
        for text in [str(o.code or "") for o in accepted] + [str(o.observation or "") for o in accepted]:
            for match in A1_REFERENCE.finditer(text):
                candidates.append((
                    match.group("quoted") or match.group("sheet") or "",
                    int(match.group("row")),
                    _column_index(match.group("col")),
                ))
            candidates.extend(("", int(m.group("row")), int(m.group("col"))) for m in CELL_REPR.finditer(text))
        anchors: list[tuple[Any, int, int]] = []
        for sheet_name, row, col in candidates:
            sheet = self._sheet(sheet_name)
            if sheet is None or row > sheet.max_row or col > sheet.max_column:
                continue
            if (sheet, row, col) not in anchors:
                anchors.append((sheet, row, col))
        sections: list[str] = []
        shown: set[tuple[str, int, int]] = set()
        used = 0
        for sheet, row, col in anchors:
            if len(sections) >= MAX_CONTEXT_ANCHORS:
                break
            if (sheet.title, row, col) in shown:
                continue
            section = _window(sheet, row, col)
            if used + len(section) > MAX_CONTEXT_CHARS:
                break
            sections.append(section)
            used += len(section)
            shown.update(
                (sheet.title, r, c)
                for r in range(row - CONTEXT_ROWS, row + CONTEXT_ROWS + 1)
                for c in range(col - CONTEXT_COLUMNS, col + CONTEXT_COLUMNS + 1)
            )
        return "\n\n".join(sections) or "The code referenced no worksheet cells by address."

    def _sheet(self, name: str) -> Any:
        operators = getattr(self.env, "operators", None)
        try:
            return operators._workbook.resolve_sheet(name)
        except Exception:
            return None

    def _grouped_header_context(self) -> str:
        """Expose verified sibling headers so final review can detect partial group coverage."""
        operators = getattr(self.env, "operators", None)
        if operators is None or not hasattr(operators, "list_headers"):
            return "No grouped headers were available to the final reviewer."
        namespace = getattr(self.env, "execution_namespace", {})
        table_ids = namespace.get("selected_table_ids") if isinstance(namespace, dict) else None
        if isinstance(table_ids, str):
            table_ids = [table_ids]
        if not table_ids:
            try:
                table_ids = operators.list_tables() if hasattr(operators, "list_tables") else []
            except Exception:
                table_ids = []

        lines = []
        for table_id in table_ids or []:
            try:
                headers = operators.list_headers(str(table_id))
            except Exception:
                continue
            for header in headers:
                children = getattr(header, "sub_headers", []) or []
                if not children:
                    continue
                child_text = ", ".join(
                    f"{getattr(child, 'id', '')} ({getattr(child, 'label', '')})"
                    for child in children
                )
                lines.append(
                    f"table={table_id}; parent={getattr(header, 'id', '')} "
                    f"({getattr(header, 'label', '')}); children=[{child_text}]"
                )
        return "\n".join(lines) if lines else "No grouped headers were available to the final reviewer."

    def _group_context(self, question: str) -> str:
        from TableAgent.stages.qa.header_hints import _contains_phrase, _normalize
        from TableAgent.utils import range_to_a1

        operators = getattr(self.env, "operators", None)
        if operators is None or not hasattr(operators, "list_groups"):
            return "No relevant structure groups."
        namespace = getattr(self.env, "execution_namespace", {})
        table_ids = namespace.get("selected_table_ids") if isinstance(namespace, dict) else None
        if isinstance(table_ids, str):
            table_ids = [table_ids]
        if not table_ids:
            table_ids = operators.list_tables()
        exact = []
        fallback = []
        normalized_question = _normalize(question)
        for table_id in table_ids or []:
            for group in operators.list_groups(str(table_id)):
                line = (
                    f"table={table_id}; group={group.id}; label={group.label}; axis={group.axis}; "
                    f"group_range={range_to_a1(group.group_range) if group.group_range else None}; "
                    f"data_range={range_to_a1(group.data_range) if group.data_range else None}"
                )
                fallback.append(line)
                if _contains_phrase(normalized_question, _normalize(group.label)):
                    exact.append(line)
        if exact:
            header = (
                "Matched to the question by label. The answer's evidence must stay inside these ranges."
            )
            lines = exact[:8]
        elif fallback:
            header = (
                "No group label matched the question. Listed as layout context only -- do not treat "
                "membership in any of these as something the answer claimed."
            )
            lines = fallback[:8]
        else:
            return "No relevant structure groups."
        return "\n".join([header, *lines])


__all__ = ["ReviewFinalAnswerAction"]


def _column_index(letters: str) -> int:
    index = 0
    for letter in letters:
        index = index * 26 + ord(letter) - ord("A") + 1
    return index


def _window(sheet: Any, row: int, col: int) -> str:
    """Column headers, the enclosing section label, and nearby rows around one cell."""
    from openpyxl.utils import get_column_letter

    label_col = next(
        (c for c in range(1, col) if isinstance(sheet.cell(row, c).value, str) and sheet.cell(row, c).value.strip()),
        1,
    )
    columns = sorted({label_col, *range(max(1, col - CONTEXT_COLUMNS), min(sheet.max_column, col + CONTEXT_COLUMNS) + 1)})
    merged = {
        (r, c): (block.min_row, block.min_col)
        for block in sheet.merged_cells.ranges
        if block.min_row < row
        for r in range(block.min_row, block.max_row + 1)
        for c in range(block.min_col, block.max_col + 1)
    }
    headers = []
    for c in columns:
        if c == label_col:
            continue
        parts = []
        for r in range(1, row):
            value = sheet.cell(*merged.get((r, c), (r, c))).value
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                break
            if isinstance(value, str) and value.strip():
                parts.append(" ".join(value.split()))
        if parts:
            headers.append(f"{get_column_letter(c)}: {' / '.join(parts[-3:])}")
    section = None
    for r in range(row - 1, max(0, row - 60), -1):
        values = [cell for cell in sheet[r] if cell.value not in (None, "")]
        if len(values) == 1 and isinstance(values[0].value, str):
            section = f"{values[0].coordinate}:{_cell_text(values[0])}"
            break
    lines = [f"### '{sheet.title}'!{get_column_letter(col)}{row}"]
    if headers:
        lines.append("column headers: " + "; ".join(headers))
    if section:
        lines.append(f"section above: {section}")
    for r in range(max(1, row - CONTEXT_ROWS), min(sheet.max_row, row + CONTEXT_ROWS) + 1):
        cells = " | ".join(f"{sheet.cell(r, c).coordinate}:{_value_text(sheet.cell(r, c))}" for c in columns)
        lines.append(f"{'>' if r == row else ' '} | {cells} |")
    return "\n".join(lines)


def _value_text(cell: Any) -> str:
    """A cell's text, with a percentage-formatted number also shown as the workbook displays it.

    A stored 1 formatted as 0.0% reads as 100%, not 1; without the format the reviewer sees only
    the stored value and can mistake a correct share for an error.
    """
    text = _cell_text(cell)
    value = cell.value
    number_format = str(getattr(cell, "number_format", "") or "")
    if isinstance(value, bool) or not isinstance(value, (int, float)) or "%" not in number_format:
        return text
    decimals = PERCENT_DECIMALS.search(number_format)
    places = len(decimals.group(1)) if decimals else 0
    return f"{text} (shown as {value * 100:.{places}f}%)"
