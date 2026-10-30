"""The `sheet` object of the notebook: reads, searches and describes workbook cells.

Reads resolve merged cells to the value a reader sees. structure.yaml is loaded as a rough map:
it names headers and groups with their cell ranges, and may be incomplete or off by a row.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

import pandas as pd
import yaml
from openpyxl.utils import get_column_letter

from TableAgent.domain.ranges import CellRange
from TableAgent.stages.qa.prompts import HEADERS_LOOK_QUESTION
from TableAgent.stages.qa.tools.headers import HeaderWalker, cell_format, format_text, headers_line, plain, text_of
from TableAgent.utils.excel_utils import cell_to_a1, parse_a1_range, range_to_a1

MAX_HEADER_LOOKS = 3  # images `headers` asks a vision model about, per question
HEADER_LOOK_ROWS = 40  # the image runs from up to this many rows above the cell ...
HEADER_LOOK_COLUMNS = 20  # ... and this many columns to its left, the cell included
MAX_FORMAT_CELLS = 300  # filled cells `formats` lists


@dataclass
class StructureEntry:
    """A header or group of structure.yaml: its label path, where its label and data sit, and what it holds."""

    kind: str  # "header" or "group"
    path: list[str]
    sheet: str
    label_range: CellRange | None
    data_range: CellRange | None
    description: str = ""
    axis: str = ""  # of a group: the dimension it groups along, e.g. "region"


@dataclass
class StructureTable:
    name: str
    sheet: str
    description: str
    entries: list[StructureEntry] = field(default_factory=list)


def load_structure(paths: list[Path]) -> list[StructureTable]:
    """Tables of every readable structure.yaml; an unreadable file contributes nothing."""
    tables: list[StructureTable] = []
    for path in paths:
        try:
            payload = yaml.safe_load(Path(path).read_text(encoding="utf-8"))
        except (OSError, yaml.YAMLError):
            continue
        if not isinstance(payload, dict):
            continue
        for key, data in payload.items():
            if key == "relations" or not isinstance(data, dict):
                continue
            sheet = str(data.get("sheet") or data.get("name") or key)
            table = StructureTable(str(data.get("name") or key), sheet, str(data.get("description") or ""))
            for header in data.get("headers") or []:
                _add_header(table, header, [], sheet)
            for group in data.get("groups") or []:
                if isinstance(group, dict):
                    table.entries.append(StructureEntry(
                        "group",
                        [str(group.get("label") or group.get("id") or "")],
                        sheet,
                        _range(group.get("group_range") or group.get("range"), sheet),
                        _range(group.get("data_range"), sheet),
                        str(group.get("description") or ""),
                        str(group.get("axis") or ""),
                    ))
            tables.append(table)
    return tables


def _add_header(table: StructureTable, header: Any, parents: list[str], sheet: str) -> None:
    if not isinstance(header, dict):
        return
    path = [*parents, str(header.get("label") or header.get("id") or "")]
    table.entries.append(StructureEntry(
        "header", path, sheet, _range(header.get("header_range"), sheet), _range(header.get("data_range"), sheet),
        str(header.get("description") or ""),
    ))
    for sub in header.get("sub_headers") or []:
        _add_header(table, sub, path, sheet)


def _range(value: Any, sheet: str) -> CellRange | None:
    if not value:
        return None
    try:
        return parse_a1_range(str(value), sheet)
    except ValueError:
        return None


class SheetTools:
    """Read, search and describe workbook cells; merged cells read as their visible value."""

    def __init__(
        self,
        workbook: Any,
        structure: list[StructureTable] | None = None,
        look: Callable[[str, str, str], str] | None = None,
        formula_path: Path | None = None,
    ):
        self.workbook = workbook
        self.structure = structure or []
        self._look = look  # (ref, question, sheet) -> the vision model's answer; None without a vision model
        self._walker = HeaderWalker(self, formula_path)
        self._header_looks = 0
        self._merged: dict[str, dict[tuple[int, int], tuple[int, int]]] = {}

    def sheets(self) -> list[str]:
        """Sheet names with their used size, e.g. `Data (233 rows x 19 columns)`."""
        return [f"{ws.title} ({ws.max_row} rows x {ws.max_column} columns)" for ws in self.workbook.worksheets]

    def worksheet(self, sheet: str = "") -> Any:
        """The named sheet, matched ignoring case and spacing; the first sheet when unnamed."""
        names = self.workbook.sheetnames
        if not sheet:
            return self.workbook[names[0]]
        wanted = plain(sheet)
        for name in names:
            if plain(name) == wanted:
                return self.workbook[name]
        if len(names) == 1:
            return self.workbook[names[0]]
        raise ValueError(f"No sheet named {sheet!r}; sheets: {names}")

    def value(self, row: int, col: int, sheet: str = "") -> Any:
        """The visible value at a position: inside a merged range, the range's top-left value."""
        ws = self.worksheet(sheet)
        anchor_row, anchor_col = self._merged_map(ws).get((row, col), (row, col))
        return ws.cell(anchor_row, anchor_col).value

    def read(self, ref: str, sheet: str = "") -> list[list[Any]]:
        """Values of an A1 cell or range (e.g. "B3:F9" or "'Data'!B3"), row by row."""
        cells = parse_a1_range(ref, sheet)
        return [
            [self.value(row, col, cells.sheet) for col in range(cells.start_col, cells.end_col + 1)]
            for row in range(cells.start_row, cells.end_row + 1)
        ]

    def find(self, text: Any, sheet: str = "") -> list[tuple[str, Any]]:
        """(A1 ref, value) of every cell whose text contains `text` as whole words, ignoring case and spacing."""
        query = _words(text)
        if not query:
            return []
        ws = self.worksheet(sheet)
        return [
            (cell.coordinate, cell.value)
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None and f" {query} " in f" {_words(cell.value)} "
        ]

    def frame(self, sheet: str = "") -> pd.DataFrame:
        """The whole used range as a DataFrame: index = Excel row numbers, columns = column letters."""
        ws = self.worksheet(sheet)
        rows = [
            [self.value(row, col, ws.title) for col in range(1, ws.max_column + 1)]
            for row in range(1, ws.max_row + 1)
        ]
        columns = [get_column_letter(col) for col in range(1, ws.max_column + 1)]
        return pd.DataFrame(rows, columns=columns, index=range(1, ws.max_row + 1))

    def headers(self, ref: str, sheet: str = "", look: bool = True) -> dict[str, Any]:
        """What one cell means: its row labels, column headers, sections, block and structure.yaml entries.

        Found by walking the layout in code (tools/headers.py). When the walk is unsure and a vision model is
        available, `vision` adds its answer about an image of the area above and left of the cell.
        """
        cells = parse_a1_range(ref, sheet)
        ws = self.worksheet(cells.sheet)
        found = self._walker.headers(ws, cells.start_row, cells.start_col)
        if look and found["uncertain"] and self._look is not None and self._header_looks < MAX_HEADER_LOOKS:
            self._header_looks += 1
            area = (
                f"{cell_to_a1(max(1, cells.start_row - HEADER_LOOK_ROWS + 1), max(1, cells.start_col - HEADER_LOOK_COLUMNS + 1))}"
                f":{found['cell']}"
            )
            try:
                found["vision"] = self._look(area, HEADERS_LOOK_QUESTION.format(cell=found["cell"]), ws.title)
            except Exception as exc:  # the image cannot be rendered or the call budget is used up
                found["vision"] = f"unavailable: {exc}"
        return found

    def describe(self, ref: str, sheet: str = "") -> str:
        """One line for `headers(ref)` without the vision model, e.g. `D8 = 35 | row: Alpha | column: 2019`."""
        return headers_line(self.headers(ref, sheet, look=False))

    def formats(self, ref: str, sheet: str = "") -> str:
        """The filled cells of a range with their format and formula, a line each, e.g. `D9 = 45 | bold | =SUM(D5:D8)`."""
        cells = parse_a1_range(ref, sheet)
        ws = self.worksheet(cells.sheet)
        formulas = self._walker.formula_sheet(ws.title)
        lines = []
        for row in range(cells.start_row, cells.end_row + 1):
            for col in range(cells.start_col, cells.end_col + 1):
                value = ws.cell(row, col).value
                if value in (None, ""):
                    continue
                if len(lines) == MAX_FORMAT_CELLS:
                    return "\n".join(lines + [f"... cut at {MAX_FORMAT_CELLS} cells"])
                parts = [f"{cell_to_a1(row, col)} = {text_of(value)}", format_text(cell_format(ws.cell(row, col)))]
                formula = formulas.cell(row, col).value if formulas is not None else None
                if isinstance(formula, str) and formula.startswith("="):
                    parts.append(formula)
                lines.append(" | ".join(parts))
        return "\n".join(lines)

    def around(self, ref: str, rows: int = 2, cols: int = 2, sheet: str = "") -> str:
        """The filled cells around one cell, a line per row with their refs; the cell itself is marked *."""
        cells = parse_a1_range(ref, sheet)
        ws = self.worksheet(cells.sheet)
        row, col = cells.start_row, cells.start_col
        lines = []
        for r in range(max(1, row - rows), row + rows + 1):
            parts = [
                f"{'*' if (r, c) == (row, col) else ''}{cell_to_a1(r, c)}={text_of(value)}"
                for c in range(max(1, col - cols), col + cols + 1)
                if (value := self.value(r, c, ws.title)) not in (None, "") or (r, c) == (row, col)
            ]
            lines.append(f"row {r}: " + (" | ".join(parts) if parts else "(empty)"))
        return "\n".join(lines)

    def look(self, ref: str, question: str, sheet: str = "") -> str:
        """Ask a vision model about an image of a cell range, for what values do not show (ticks, shading, spans)."""
        if self._look is None:
            raise ValueError("sheet.look is not available in this run.")
        return self._look(ref, question, self.worksheet(sheet).title)

    def structure_at(self, row: int, col: int, sheet: str = "", label: bool = False) -> list[str]:
        """structure.yaml headers and groups whose data range (or label range) holds a cell, as `header: A > B`."""
        return [f"{entry.kind}: " + " > ".join(entry.path) for entry in self.structure_entries_at(row, col, sheet, label)]

    def structure_entries_at(self, row: int, col: int, sheet: str = "", label: bool = False) -> list[StructureEntry]:
        """The structure.yaml entries whose data range (or label range) holds a cell."""
        ws_title = self.worksheet(sheet).title
        found = []
        for table in self.structure:
            for entry in table.entries:
                cells = entry.label_range if label else entry.data_range
                if cells is None or plain(entry.sheet) != plain(ws_title):
                    continue
                if cells.start_row <= row <= cells.end_row and cells.start_col <= col <= cells.end_col:
                    found.append(entry)
        return found

    def structure_find(self, text: str, sheet: str = "") -> list[dict[str, Any]]:
        """structure.yaml headers and groups whose label path, or else description, holds `text` as whole words.

        Entries matched by label come first; each gives its path, label and data ranges, and description.
        """
        query = _words(text)
        found: dict[str, list[dict[str, Any]]] = {"label": [], "description": []}
        for table in self.structure:
            for entry in table.entries:
                if not query or (sheet and plain(entry.sheet) != plain(sheet)):
                    continue
                matched = (
                    "label" if f" {query} " in f" {_words(' '.join(entry.path))} "
                    else "description" if f" {query} " in f" {_words(entry.description)} " else ""
                )
                if matched:
                    found[matched].append({
                        "kind": entry.kind,
                        "path": " > ".join(entry.path),
                        "sheet": entry.sheet,
                        "label_range": range_to_a1(entry.label_range) if entry.label_range else None,
                        "data_range": range_to_a1(entry.data_range) if entry.data_range else None,
                        "description": entry.description,
                        **({"axis": entry.axis} if entry.axis else {}),
                        "matched": matched,
                    })
        return found["label"] + found["description"]

    def _merged_map(self, ws: Any) -> dict[tuple[int, int], tuple[int, int]]:
        if ws.title not in self._merged:
            self._merged[ws.title] = {
                (r, c): (merged.min_row, merged.min_col)
                for merged in ws.merged_cells.ranges
                for r in range(merged.min_row, merged.max_row + 1)
                for c in range(merged.min_col, merged.max_col + 1)
                if (r, c) != (merged.min_row, merged.min_col)
            }
        return self._merged[ws.title]


def _words(value: Any) -> str:
    return " ".join(re.findall(r"\w+", unicodedata.normalize("NFKC", str(value)).casefold()))
