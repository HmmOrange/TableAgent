from pathlib import Path
from typing import Any

import openpyxl
import yaml
from openpyxl.utils import get_column_letter

from TableAgent.rendering.workbook import WorkbookRenderer

from .contracts import UnderstandingInput, UnderstandingOutput
from .parsing import parse_understanding
from .prompts import REPAIR_PROMPT_TEMPLATE, SYSTEM_PROMPT, USER_PROMPT_TEMPLATE

class UnderstandingStage:
    def __init__(self, renderer: WorkbookRenderer, vlm: Any):
        self.renderer = renderer
        self.vlm = vlm

    def run(self, stage_input: UnderstandingInput) -> UnderstandingOutput:
        workbook_path = Path(stage_input.workbook_path)
        artifact_dir = Path(stage_input.artifact_dir)
        artifact_dir.mkdir(parents=True, exist_ok=True)
        viewport_range = _used_range(workbook_path, stage_input.sheet_name)
        image_path = artifact_dir / "worksheet.png"
        self.renderer.source_viewport_to_image(
            workbook_path,
            stage_input.sheet_name,
            viewport_range,
            image_path,
        )
        prompt = USER_PROMPT_TEMPLATE.format(
            workbook_name=workbook_path.name,
            sheet_name=stage_input.sheet_name,
            viewport_range=viewport_range,
        )
        (artifact_dir / "prompt.txt").write_text(prompt, encoding="utf-8")
        response = self.vlm.generate_with_image(
            prompt=prompt,
            image_path=image_path,
            system_prompt=SYSTEM_PROMPT,
        )
        (artifact_dir / "response.txt").write_text(response.content, encoding="utf-8")
        try:
            understanding = parse_understanding(response.content)
        except ValueError as first_error:
            repair_prompt = REPAIR_PROMPT_TEMPLATE.format(
                error=str(first_error),
                response=response.content,
            )
            repair_response = self.vlm.generate_with_image(
                prompt=repair_prompt,
                image_path=image_path,
                system_prompt=SYSTEM_PROMPT,
            )
            (artifact_dir / "repair_prompt.txt").write_text(repair_prompt, encoding="utf-8")
            (artifact_dir / "repair_response.txt").write_text(
                repair_response.content,
                encoding="utf-8",
            )
            understanding = parse_understanding(repair_response.content)
            response = repair_response

        result_path = artifact_dir / "understanding.yaml"
        result_path.write_text(
            yaml.safe_dump(
                understanding.to_dict(),
                sort_keys=False,
                allow_unicode=True,
            ),
            encoding="utf-8",
        )
        return UnderstandingOutput(
            understanding=understanding,
            viewport_range=viewport_range,
            image_path=image_path,
            result_path=result_path,
            response=response,
        )


def _used_range(workbook_path: Path, sheet_name: str) -> str:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {sheet_name}")
        worksheet = workbook[sheet_name]
        used_range = f"A1:{get_column_letter(max(1, worksheet.max_column))}{max(1, worksheet.max_row)}"
    finally:
        workbook.close()
    return used_range
