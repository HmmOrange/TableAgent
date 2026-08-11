import shutil
import struct
import subprocess
import tempfile
import zlib
from pathlib import Path

import pymupdf
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

from .helper import SheetHelper

class SheetImager:
    @staticmethod
    def _resolve_soffice_path(soffice_path):
        if shutil.which(soffice_path):
            return soffice_path

        default_macos_path = "/opt/homebrew/bin/soffice"
        if Path(default_macos_path).exists():
            return default_macos_path

        return soffice_path

    @staticmethod
    def _png_chunk(chunk_type, chunk_data):
        chunk = chunk_type + chunk_data
        return (
            struct.pack(">I", len(chunk_data))
            + chunk
            + struct.pack(">I", zlib.crc32(chunk) & 0xFFFFFFFF)
        )

    @staticmethod
    def _unfilter_png_scanlines(raw_data, width, height, bpp):
        stride = width * bpp
        rows = []
        prev_row = bytearray(stride)
        pos = 0

        for _ in range(height):
            filter_type = raw_data[pos]
            pos += 1

            row = bytearray(raw_data[pos : pos + stride])
            pos += stride

            for i in range(stride):
                left = row[i - bpp] if i >= bpp else 0
                up = prev_row[i]
                up_left = prev_row[i - bpp] if i >= bpp else 0

                if filter_type == 1:
                    row[i] = (row[i] + left) & 0xFF
                elif filter_type == 2:
                    row[i] = (row[i] + up) & 0xFF
                elif filter_type == 3:
                    row[i] = (row[i] + ((left + up) // 2)) & 0xFF
                elif filter_type == 4:
                    p = left + up - up_left
                    pa = abs(p - left)
                    pb = abs(p - up)
                    pc = abs(p - up_left)

                    predictor = (
                        left if pa <= pb and pa <= pc
                        else up if pb <= pc
                        else up_left
                    )

                    row[i] = (row[i] + predictor) & 0xFF

            rows.append(row)
            prev_row = row

        return rows

    @staticmethod
    def _crop_png_to_visible_area(img_path):
        img_path = Path(img_path)

        data = img_path.read_bytes()

        if data[:8] != b"\x89PNG\r\n\x1a\n":
            return

        pos = 8
        idat_data = b""

        while pos < len(data):
            chunk_len = struct.unpack(">I", data[pos : pos + 4])[0]
            chunk_type = data[pos + 4 : pos + 8]
            chunk_data = data[pos + 8 : pos + 8 + chunk_len]
            pos += chunk_len + 12

            if chunk_type == b"IHDR":
                width, height, bit_depth, color_type, _, _, interlace = (
                    struct.unpack(">IIBBBBB", chunk_data)
                )
            elif chunk_type == b"IDAT":
                idat_data += chunk_data

        if bit_depth != 8 or color_type != 6 or interlace != 0:
            return

        bpp = 4

        rows = SheetImager._unfilter_png_scanlines(
            zlib.decompress(idat_data),
            width,
            height,
            bpp,
        )

        min_x, min_y = width, height
        max_x, max_y = -1, -1

        for y, row in enumerate(rows):
            for x in range(width):
                alpha = row[x * bpp + 3]

                if alpha:
                    min_x = min(min_x, x)
                    min_y = min(min_y, y)
                    max_x = max(max_x, x)
                    max_y = max(max_y, y)

        if max_x == -1:
            return

        new_width = max_x - min_x + 1
        new_height = max_y - min_y + 1

        cropped_rows = []

        for row in rows[min_y : max_y + 1]:
            start = min_x * bpp
            end = (max_x + 1) * bpp
            cropped_rows.append(b"\x00" + bytes(row[start:end]))

        png_data = (
            b"\x89PNG\r\n\x1a\n"
            + SheetImager._png_chunk(
                b"IHDR",
                struct.pack(
                    ">IIBBBBB",
                    new_width,
                    new_height,
                    8,
                    6,
                    0,
                    0,
                    0,
                ),
            )
            + SheetImager._png_chunk(
                b"IDAT",
                zlib.compress(b"".join(cropped_rows)),
            )
            + SheetImager._png_chunk(b"IEND", b"")
        )

        img_path.write_bytes(png_data)

    @staticmethod
    def _pdf_to_png(pdf_path, img_path, dpi=200):
        doc = pymupdf.open(pdf_path)

        try:
            page = doc[0]
            pix = page.get_pixmap(dpi=dpi, alpha=True)
            pix.save(img_path)
        finally:
            doc.close()

    @staticmethod
    def to_image(
        file_path: str,
        sheet_name: str,
        output_path: str,
        dpi: int = 200,
        soffice_path: str = "soffice",
    ):
        file_path = Path(file_path)
        output_path = Path(output_path)

        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = load_workbook(file_path)

        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found")

        ws = wb[sheet_name]

        nrows, ncols = SheetHelper.get_n_row_col(ws)

        if nrows == 0 or ncols == 0:
            return None

        old_active_idx = wb.index(wb.active)
        old_states = {
            s.title: s.sheet_state
            for s in wb.worksheets
        }

        old_print_area = ws.print_area
        old_fit_to_page = ws.sheet_properties.pageSetUpPr.fitToPage
        old_fit_to_width = ws.page_setup.fitToWidth
        old_fit_to_height = ws.page_setup.fitToHeight

        old_margins = (
            ws.page_margins.left,
            ws.page_margins.right,
            ws.page_margins.top,
            ws.page_margins.bottom,
        )

        try:
            last_col = get_column_letter(ncols)

            ws.print_area = f"A1:{last_col}{nrows}"

            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 1

            ws.page_margins.left = 0
            ws.page_margins.right = 0
            ws.page_margins.top = 0
            ws.page_margins.bottom = 0

            for s in wb.worksheets:
                s.sheet_state = (
                    "visible" if s.title == sheet_name else "hidden"
                )

            wb.active = wb.index(ws)

            with tempfile.TemporaryDirectory() as tmp:
                tmp = Path(tmp)

                xlsx_path = tmp / "sheet.xlsx"
                pdf_path = tmp / "sheet.pdf"
                profile_path = tmp / "libreoffice_profile"

                profile_path.mkdir()

                wb.save(xlsx_path)

                result = subprocess.run(
                    [
                        SheetImager._resolve_soffice_path(soffice_path),
                        f"-env:UserInstallation={profile_path.resolve().as_uri()}",
                        "--headless",
                        "--convert-to",
                        "pdf",
                        "--outdir",
                        str(tmp),
                        str(xlsx_path),
                    ],
                    check=True,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                )

                if not pdf_path.exists():
                    raise RuntimeError(
                        "LibreOffice did not create PDF file.\n"
                        f"stdout:\n{result.stdout}\n"
                        f"stderr:\n{result.stderr}"
                    )

                SheetImager._pdf_to_png(
                    pdf_path,
                    tmp / "sheet.png",
                    dpi=dpi,
                )

                rendered_path = tmp / "sheet.png"
                SheetImager._crop_png_to_visible_area(rendered_path)

                if output_path.suffix.lower() == ".png":
                    shutil.copyfile(rendered_path, output_path)
                else:
                    with Image.open(rendered_path) as image:
                        # JPEG and several other output formats do not support
                        # the alpha channel used to crop the rendered PDF.
                        rgb_image = Image.new("RGB", image.size, "white")
                        if image.mode in ("RGBA", "LA"):
                            rgb_image.paste(image, mask=image.getchannel("A"))
                        else:
                            rgb_image.paste(image)
                        rgb_image.save(output_path)

        finally:
            ws.print_area = old_print_area

            ws.sheet_properties.pageSetUpPr.fitToPage = old_fit_to_page
            ws.page_setup.fitToWidth = old_fit_to_width
            ws.page_setup.fitToHeight = old_fit_to_height

            (
                ws.page_margins.left,
                ws.page_margins.right,
                ws.page_margins.top,
                ws.page_margins.bottom,
            ) = old_margins

            for s in wb.worksheets:
                s.sheet_state = old_states[s.title]

            wb.active = old_active_idx

        return str(output_path)
    
