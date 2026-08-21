from __future__ import annotations

from dataclasses import dataclass
from copy import copy
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment


@dataclass(frozen=True)
class CompressionConfig:
    similarity_threshold: float = 0.8
    keep_rows: int = 2
    insert_ellipsis: bool = False

    def __post_init__(self) -> None:
        if not 0 < self.similarity_threshold <= 1:
            raise ValueError("compression similarity threshold must be in (0, 1]")
        if self.keep_rows < 0:
            raise ValueError("compression keep_rows must be non-negative")


@dataclass(frozen=True)
class RowMap:
    compressed_row: int
    original_start: int
    original_end: int
    marker: bool = False

    def to_dict(self) -> dict[str, Any]:
        return {
            "compressed_row": self.compressed_row,
            "original_start": self.original_start,
            "original_end": self.original_end,
            "marker": self.marker,
        }


@dataclass(frozen=True)
class SheetCompressionResult:
    sheet_name: str
    original_rows: int
    compressed_rows: int
    rows: tuple[RowMap, ...]

    def to_dict(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet_name,
            "original_rows": self.original_rows,
            "compressed_rows": self.compressed_rows,
            "rows": [row.to_dict() for row in self.rows],
        }

    def map_row(self, row: int) -> tuple[int, int]:
        if not self.rows:
            return row, row
        item = self.rows[min(max(row, 1), len(self.rows)) - 1]
        return item.original_start, item.original_end


class SheetCompressor:
    """Compress consecutive, style/type-similar rows in an XLSX workbook."""

    def __init__(self, config: CompressionConfig | None = None, **kwargs: Any):
        self.config = config or CompressionConfig(
            similarity_threshold=float(kwargs.get("similarity_threshold", 0.8)),
            keep_rows=int(kwargs.get("keep_rows", 2)),
            insert_ellipsis=bool(kwargs.get("insert_ellipsis", False)),
        )

    @staticmethod
    def _cell_type(value: Any) -> str:
        if value is None or (isinstance(value, str) and not value.strip()):
            return "empty"
        if isinstance(value, bool):
            return "boolean"
        if isinstance(value, (int, float)):
            return "numeric"
        if isinstance(value, str) and value.startswith("="):
            return "formula"
        return "text"

    @staticmethod
    def _cell_fp(cell: Any) -> tuple[Any, ...]:
        color = getattr(cell.font.color, "rgb", None) if cell.font.color else None
        fill = getattr(cell.fill.fgColor, "rgb", None) if cell.fill else None
        return (
            SheetCompressor._cell_type(cell.value),
            cell.font.bold,
            cell.font.italic,
            cell.font.underline,
            color,
            fill,
            cell.number_format,
        )

    def _row_fp(self, worksheet: Any, row: int, max_col: int) -> tuple[tuple[Any, ...], ...]:
        return tuple(self._cell_fp(worksheet.cell(row, col)) for col in range(1, max_col + 1))

    def _similar(self, left: tuple[tuple[Any, ...], ...], right: tuple[tuple[Any, ...], ...]) -> bool:
        if not left:
            return True
        matches = sum(a == b for a, b in zip(left, right))
        return matches / len(left) >= self.config.similarity_threshold

    @staticmethod
    def _copy_row(worksheet: Any, source: int, target: int, max_col: int) -> None:
        for col in range(1, max_col + 1):
            source_cell = worksheet.cell(source, col)
            target_cell = worksheet.cell(target, col)
            if isinstance(target_cell, MergedCell):
                continue
            target_cell.value = source_cell.value
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            target_cell.number_format = source_cell.number_format

    def _marker(self, worksheet: Any, row: int, max_col: int, source: int) -> None:
        for col in range(1, max_col + 1):
            cell = worksheet.cell(row, col)
            if isinstance(cell, MergedCell):
                continue
            source_cell = worksheet.cell(source, col)
            cell._style = copy(source_cell._style)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.value = "..."

    def compress_sheet(self, worksheet: Any) -> SheetCompressionResult:
        original_rows = worksheet.max_row
        max_col = max(worksheet.max_column, 1)
        fingerprints = [None] + [self._row_fp(worksheet, row, max_col) for row in range(1, original_rows + 1)]
        output: list[tuple[int, int, int, bool]] = []
        row = 1
        while row <= original_rows:
            end = row
            while end < original_rows and self._similar(fingerprints[end], fingerprints[end + 1]):
                end += 1
            run_length = end - row + 1
            if run_length > self.config.keep_rows * 2:
                for original in range(row, row + self.config.keep_rows):
                    output.append((original, original, original, False))
                if self.config.insert_ellipsis:
                    output.append((row + self.config.keep_rows, row + self.config.keep_rows, end - self.config.keep_rows, True))
                for original in range(end - self.config.keep_rows + 1, end + 1):
                    output.append((original, original, original, False))
            else:
                output.extend((original, original, original, False) for original in range(row, end + 1))
            row = end + 1

        # Rebuild rows from the original snapshot so deletion is deterministic.
        values = []
        for compressed, start, end, marker in output:
            if marker:
                values.append((compressed, start, end, marker, None))
            else:
                values.append((compressed, start, end, marker, [worksheet.cell(start, col).value for col in range(1, max_col + 1)]))
        worksheet.delete_rows(1, worksheet.max_row)
        for target, (_, start, end, marker, row_values) in enumerate(values, start=1):
            worksheet.append([None] * max_col)
            if marker:
                self._marker(worksheet, target, max_col, max(start, 1))
            else:
                for col, value in enumerate(row_values or [], start=1):
                    worksheet.cell(target, col).value = value

        # openpyxl keeps stale cells and merged ranges after delete_rows().
        # Remove anything outside the rebuilt compressed rectangle so layout
        # metadata and viewport traversal cannot see original tail rows.
        final_rows = len(values)
        for merged_range in list(worksheet.merged_cells.ranges):
            if merged_range.max_row > final_rows or merged_range.max_col > max_col:
                worksheet.merged_cells.ranges.remove(merged_range)
        for coordinate in list(worksheet._cells):
            if coordinate[0] > final_rows or coordinate[1] > max_col:
                del worksheet._cells[coordinate]

        mapping = tuple(RowMap(index, start, end, marker) for index, (_, start, end, marker, _) in enumerate(values, start=1))
        return SheetCompressionResult(worksheet.title, original_rows, len(mapping), mapping)

    def compress_workbook(self, input_path: Path, output_path: Path, mapping_path: Path, sheets: tuple[str, ...] = ()) -> dict[str, SheetCompressionResult]:
        workbook = load_workbook(input_path)
        try:
            selected = set(sheets)
            results = {
                worksheet.title: self.compress_sheet(worksheet)
                for worksheet in workbook.worksheets
                if not selected or worksheet.title in selected
            }
            output_path.parent.mkdir(parents=True, exist_ok=True)
            mapping_path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(output_path)
            mapping_path.write_text(json.dumps({name: result.to_dict() for name, result in results.items()}, indent=2), encoding="utf-8")
            return results
        finally:
            workbook.close()
