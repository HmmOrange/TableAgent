from openpyxl.worksheet.worksheet import Worksheet

class SheetHelper:
    @staticmethod
    def get_n_row_col(ws: Worksheet):
        last_row = 0
        last_col = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    last_row = max(last_row, cell.row)
                    last_col = max(last_col, cell.column)
        return last_row, last_col

    @staticmethod
    def delete_rows(ws: Worksheet, start_row_idx: int, row_count: int = 1):
        end_row_idx = start_row_idx + row_count - 1

        old_ranges = list(ws.merged_cells.ranges)
        ws.merged_cells.ranges = []

        new_ranges = []

        for rng in old_ranges:
            min_r, max_r = rng.min_row, rng.max_row
            min_c, max_c = rng.min_col, rng.max_col

            # merge hoàn toàn phía trên vùng xóa
            if max_r < start_row_idx:
                new_ranges.append((min_r, max_r, min_c, max_c))

            # merge hoàn toàn phía dưới vùng xóa
            elif min_r > end_row_idx:
                new_ranges.append((min_r - row_count, max_r - row_count, min_c, max_c))
            else:
                # có giao với vùng xóa
                remaining_top = max(0, start_row_idx - min_r)
                remaining_bottom = max(0, max_r - end_row_idx)

                new_height = remaining_top + remaining_bottom
                if new_height >= 2:
                    new_min_r = min_r
                    new_max_r = min_r + new_height - 1
                    if max_r > end_row_idx:
                        shift = row_count
                        new_max_r -= shift
                    new_ranges.append((new_min_r, new_max_r, min_c, max_c))

        ws.delete_rows(start_row_idx, row_count)

        for min_r, max_r, min_c, max_c in new_ranges:
            if max_r > min_r or max_c > min_c:
                ws.merge_cells(
                    start_row=min_r,
                    end_row=max_r,
                    start_column=min_c,
                    end_column=max_c,
                )

    @staticmethod
    def insert_rows(ws: Worksheet, row_idx: int, row_count: int = 1):
        old_ranges = list(ws.merged_cells.ranges)
        ws.merged_cells.ranges = []

        new_ranges = []

        for rng in old_ranges:
            min_r, max_r = rng.min_row, rng.max_row
            min_c, max_c = rng.min_col, rng.max_col

            # chèn phía trên merge
            if row_idx < min_r:
                new_ranges.append((min_r + row_count, max_r + row_count, min_c, max_c))

            # chèn bên trong merge (bao gồm ngay trước max_r)
            elif min_r <= row_idx <= max_r:
                new_ranges.append((min_r, max_r + row_count, min_c, max_c))
            else:
                new_ranges.append((min_r, max_r, min_c, max_c))

        ws.insert_rows(row_idx, row_count)

        for min_r, max_r, min_c, max_c in new_ranges:
            ws.merge_cells(
                start_row=min_r,
                end_row=max_r,
                start_column=min_c,
                end_column=max_c,
            )
            
