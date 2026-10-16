import json
import re
import threading
from pathlib import Path
from types import SimpleNamespace

import pytest
import yaml
from openpyxl import Workbook
from openpyxl.styles import PatternFill
import TableAgent.stages.query_scoped.pipeline as query_pipeline_module

from TableAgent.configs import TableAgentConfig, load_config
from TableAgent.llm import LLMResponse
from TableAgent.pipeline import TableAgentPipeline
from TableAgent.pipeline.sample import EvalSample
from TableAgent.stages.query_scoped.pipeline import QueryScopedStructurePipeline
from TableAgent.stages.structure.verification import (
    DeterministicVerificationResult,
    DeterministicVerifier,
)
from TableAgent.stages.structure.verification.checks import verify_structure


class QueryScopedLLM:
    model_name = "query-scoped-test"
    temperature = 0.0

    def __init__(self):
        self.calls: list[tuple[str, str | None]] = []

    def generate(self, prompt: str, system_prompt: str | None = None) -> LLMResponse:
        self.calls.append((prompt, system_prompt))
        if "Summarize this spreadsheet region" in prompt:
            if "| A4 | 2025 |" in prompt:
                content = "Region contains 2024 and 2025 revenue; 2025 is at A4:B4."
            else:
                content = "Region contains Year and Revenue headers plus 2023 revenue."
        elif "Create a detailed retrieval-oriented overview of this worksheet" in prompt:
            content = "Revenue is organized by year in A1:B4; later years are in R3-R4."
        elif "Create a detailed retrieval-oriented overview of this workbook" in prompt:
            content = "Data sheet contains annual revenue in A1:B4."
        elif system_prompt and "locate spreadsheet regions" in system_prompt:
            match = re.search(r"region_id: ([^\n]*R3-R4)", prompt)
            assert match is not None
            content = json.dumps(
                {
                    "selected_regions": [
                        {
                            "region_id": match.group(1),
                            "reason": "Contains the requested 2025 row.",
                        }
                    ]
                }
            )
        elif system_prompt and "compact TableAgent structure" in system_prompt:
            content = """structure:
  revenue:
    id: revenue
    name: Annual revenue
    description: Revenue by year needed for the question
    sheet: Data
    headers:
      - id: year
        label: Year
        description: Calendar year lookup key
        orientation: column
        header_range: A1
        data_range: A2:A4
        sub_headers: []
      - id: revenue
        label: Revenue
        description: Annual revenue value
        orientation: column
        header_range: B1
        data_range: B2:B4
        sub_headers: []
"""
        else:
            raise AssertionError(f"Unexpected prompt: {prompt[:120]}")
        return LLMResponse(content=content, prompt_tokens=7, completion_tokens=3)


class InProcessVerifier(DeterministicVerifier):
    def _execute(self, workbook_path: Path, sheet_name: str, structure_path: Path):
        return verify_structure(workbook_path, sheet_name, structure_path)


class AlwaysRejectVerifier:
    def __init__(self):
        self.calls = 0

    def run(self, *, structure_text: str, **_kwargs) -> DeterministicVerificationResult:
        self.calls += 1
        rejected = yaml.safe_load(structure_text)
        for table in rejected.values():
            for header in table.get("headers", []):
                header["header_range"] = None
                header["data_range"] = None
        return DeterministicVerificationResult(
            status="not_good",
            feedback=f"Rejected attempt {self.calls}.",
            null_fields=[],
            report={"status": "not_good"},
            structure_text=yaml.safe_dump(rejected, sort_keys=False),
        )


class ChangingQueryScopedLLM(QueryScopedLLM):
    def __init__(self):
        super().__init__()
        self.structure_calls = 0

    def generate(self, prompt: str, system_prompt: str | None = None) -> LLMResponse:
        response = super().generate(prompt, system_prompt)
        if system_prompt and "compact TableAgent structure" in system_prompt:
            self.structure_calls += 1
            response.content = response.content.replace(
                "name: Annual revenue",
                f"name: Annual revenue attempt {self.structure_calls}",
            )
        return response


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Year", "Revenue"])
    sheet.append([2023, 100])
    sheet.append([2024, 150])
    sheet.append([2025, 200])
    sheet["A1000"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    workbook.save(path)


def _config(tmp_path: Path) -> dict:
    config = load_config(Path(__file__).resolve().parents[1] / "config.example.yaml")
    config["table_agent"].update(
        {
            "phase": "all",
            "artifact_dir": str(tmp_path / "outputs" / "repeat_1"),
            "run_artifact_dir": str(tmp_path / "outputs"),
            "query_scoped_structure_enabled": True,
            "query_static_cache_dir": str(tmp_path / "cache"),
            "query_static_chunk_rows": 2,
            "query_static_chunk_overlap_rows": 1,
            "query_static_header_rows": 1,
            "query_static_error_retries": 2,
            "query_static_error_retry_delay_seconds": 0,
            "query_location_max_regions": 1,
            "query_structure_max_retries": 0,
            "routing": {"retrieval": {"mode": "off"}},
        }
    )
    return config


def _sample(workbook_path: Path, sample_id: str = "question-1") -> EvalSample:
    return EvalSample(
        index=0,
        sample_id=sample_id,
        table_id="revenue",
        table_content="",
        question="What is revenue in 2025?",
        answer=["200"],
        table_path=str(workbook_path),
        raw={"selected_sheets": ["Data"]},
    )


def test_query_scoped_flow_reuses_static_cache_and_writes_per_query_structure(
    tmp_path: Path,
):
    workbook_path = tmp_path / "revenue.xlsx"
    _workbook(workbook_path)
    sample = _sample(workbook_path)
    llm = QueryScopedLLM()
    pipeline = TableAgentPipeline(llm, None, _config(tmp_path))
    pipeline.verifier = InProcessVerifier()
    pipeline._run_verified_qa = lambda **kwargs: (
        LLMResponse(content="200", prompt_tokens=11, completion_tokens=1),
        {"success": True, "execution_time": 0.01},
    )

    pipeline.prepare_samples([sample])
    assert len(llm.calls) == 5
    chunk_prompts = [
        prompt
        for prompt, _ in llm.calls
        if "Summarize this spreadsheet region" in prompt
    ]
    assert "Summary of the immediately preceding region" in chunk_prompts[1]
    pipeline.set_run_id(2)
    output = pipeline.run(sample)

    structure_path = Path(output.metadata["structure_path"])
    structure = yaml.safe_load(structure_path.read_text(encoding="utf-8"))
    assert [header["label"] for header in structure["revenue"]["headers"]] == [
        "Year",
        "Revenue",
    ]
    assert output.predicted_answer == "200"
    assert output.metadata["query_scoped_structure"] is True
    assert output.metadata["verification"]["Data"]["status"] == "good"
    assert output.metadata["location"]["selected_regions"][0]["cell_range"] == "A3:B4"
    cache_dir = Path(output.metadata["static_cache_dirs"][0])
    manifest = json.loads((cache_dir / "index.json").read_text(encoding="utf-8"))
    assert manifest["sheets"][0]["reported_max_row"] == 1000
    assert manifest["sheets"][0]["effective_max_row"] == 4

    second_llm = QueryScopedLLM()
    second_pipeline = TableAgentPipeline(second_llm, None, _config(tmp_path))
    second_pipeline.prepare_samples([sample])
    assert second_llm.calls == []


def test_query_static_reuses_complete_cache_when_only_chunk_size_changes(
    tmp_path: Path,
):
    workbook_path = tmp_path / "revenue.xlsx"
    _workbook(workbook_path)
    sample = _sample(workbook_path)

    first_pipeline = TableAgentPipeline(QueryScopedLLM(), None, _config(tmp_path))
    first_pipeline.prepare_samples([sample])
    original_indexes = list((tmp_path / "cache").rglob("index.json"))
    assert len(original_indexes) == 1

    class MustNotRunLLM(QueryScopedLLM):
        def generate(
            self,
            prompt: str,
            system_prompt: str | None = None,
        ) -> LLMResponse:
            raise AssertionError("compatible cache should have been reused")

    changed_config = _config(tmp_path)
    changed_config["table_agent"]["query_static_chunk_rows"] = 3
    second_llm = MustNotRunLLM()
    second_pipeline = TableAgentPipeline(second_llm, None, changed_config)
    second_pipeline.prepare_samples([sample])

    assert second_llm.calls == []
    assert list((tmp_path / "cache").rglob("index.json")) == original_indexes


def test_query_chunk_overlap_must_be_smaller_than_chunk_size(tmp_path: Path):
    config = _config(tmp_path)
    config["table_agent"]["query_static_chunk_overlap_rows"] = 2
    with pytest.raises(ValueError, match="must be smaller"):
        TableAgentConfig.from_config(config)


def test_query_static_ingestion_raises_instead_of_caching_llm_error(
    tmp_path: Path,
):
    workbook_path = tmp_path / "revenue.xlsx"
    _workbook(workbook_path)
    sample = _sample(workbook_path)

    class ErrorLLM(QueryScopedLLM):
        def __init__(self):
            super().__init__()
            self.attempts = 0

        def generate(
            self,
            prompt: str,
            system_prompt: str | None = None,
        ) -> LLMResponse:
            self.attempts += 1
            return LLMResponse(content="ERROR: request timed out")

    llm = ErrorLLM()
    pipeline = TableAgentPipeline(llm, None, _config(tmp_path))

    with pytest.raises(
        RuntimeError,
        match=r"LLM failed while generating chunk summary.*ERROR: request timed out",
    ):
        pipeline.prepare_samples([sample])

    assert llm.attempts == 3
    assert not list((tmp_path / "cache").rglob("index.json"))


def test_query_static_retries_only_the_failed_chunk_and_continues(
    tmp_path: Path,
):
    workbook_path = tmp_path / "revenue.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Year", "Revenue"])
    for year in range(2020, 2026):
        sheet.append([year, year - 1900])
    workbook.save(workbook_path)
    workbook.close()

    class RetryFifthChunkLLM(QueryScopedLLM):
        def __init__(self):
            super().__init__()
            self.chunk_ranges: list[str] = []
            self.failed_once = False

        def generate(
            self,
            prompt: str,
            system_prompt: str | None = None,
        ) -> LLMResponse:
            if "Summarize this spreadsheet region" in prompt:
                match = re.search(r"Region: [^\n]+ \((A\d+:B\d+)\)", prompt)
                assert match is not None
                cell_range = match.group(1)
                self.chunk_ranges.append(cell_range)
                if cell_range == "A5:B6" and not self.failed_once:
                    self.failed_once = True
                    return LLMResponse(content="ERROR: transient timeout")
            return super().generate(prompt, system_prompt)

    llm = RetryFifthChunkLLM()
    pipeline = TableAgentPipeline(llm, None, _config(tmp_path))
    pipeline.prepare_samples([_sample(workbook_path)])

    assert llm.chunk_ranges == [
        "A1:B2",
        "A2:B3",
        "A3:B4",
        "A4:B5",
        "A5:B6",
        "A5:B6",
        "A6:B7",
    ]
    summaries = sorted((tmp_path / "cache").rglob("chunk_*_summary.txt"))
    assert len(summaries) == 6
    assert all(
        not path.read_text(encoding="utf-8").startswith("ERROR:")
        for path in summaries
    )


def test_query_static_ingestion_reads_cached_values_not_formula_source(
    tmp_path: Path,
):
    workbook_path = tmp_path / "formula-heavy.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Input", "Calculated"])
    sheet.append([10, "=SUM(A2:A1000)+SUM(A2:A1000)+SUM(A2:A1000)"])
    workbook.save(workbook_path)
    workbook.close()

    pipeline = TableAgentPipeline(QueryScopedLLM(), None, _config(tmp_path))
    pipeline.prepare_samples([_sample(workbook_path)])

    source = next((tmp_path / "cache").rglob("chunk_0001_source.md"))
    source_text = source.read_text(encoding="utf-8")
    assert "=SUM(" not in source_text
    assert "| A2 | 10 |" in source_text


@pytest.mark.parametrize(
    "summary_glob",
    [
        "*/*/summary.txt",
        "*/*/sheets/Data/summary.txt",
        "*/*/sheets/Data/chunks/chunk_0001_summary.txt",
    ],
)
def test_query_static_ingestion_rebuilds_cache_containing_error_summary(
    tmp_path: Path,
    summary_glob: str,
):
    workbook_path = tmp_path / "revenue.xlsx"
    _workbook(workbook_path)
    sample = _sample(workbook_path)

    first_llm = QueryScopedLLM()
    first_pipeline = TableAgentPipeline(first_llm, None, _config(tmp_path))
    first_pipeline.prepare_samples([sample])
    failed_summary = next((tmp_path / "cache").glob(summary_glob))
    failed_summary.write_text("ERROR: stale timeout\n", encoding="utf-8")

    second_llm = QueryScopedLLM()
    second_pipeline = TableAgentPipeline(second_llm, None, _config(tmp_path))
    second_pipeline.prepare_samples([sample])

    assert second_llm.calls
    assert not failed_summary.read_text(encoding="utf-8").startswith("ERROR:")


def test_query_static_indexes_are_prepared_concurrently_across_workbooks(
    tmp_path: Path,
):
    barrier = threading.Barrier(2)
    worker_threads: set[int] = set()
    runtime = SimpleNamespace(
        settings=SimpleNamespace(max_workers=2),
        llm=object(),
        _progress=lambda _stage, **_fields: None,
    )
    pipeline = QueryScopedStructurePipeline(runtime)

    def prepare_index(path: Path, _selected_sheets: tuple[str, ...]):
        worker_threads.add(threading.get_ident())
        barrier.wait(timeout=5)
        return SimpleNamespace(
            sheets=(SimpleNamespace(sheet_name="Data"),),
            cache_hit=False,
        )

    pipeline._prepare_index = prepare_index
    samples = [
        EvalSample(0, "q-a", "a", "", "q", [], table_path=str(tmp_path / "a.xlsx")),
        EvalSample(1, "q-b", "b", "", "q", [], table_path=str(tmp_path / "b.xlsx")),
    ]
    pipeline.prepare_samples(samples)
    assert len(worker_threads) == 2


def test_query_static_reports_successes_after_another_future_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    progress_events: list[tuple[str, dict]] = []
    runtime = SimpleNamespace(
        settings=SimpleNamespace(max_workers=2),
        llm=object(),
        _progress=lambda stage, **fields: progress_events.append((stage, fields)),
    )
    pipeline = QueryScopedStructurePipeline(runtime)

    def prepare_index(path: Path, _selected_sheets: tuple[str, ...]):
        if path.name == "broken.xlsx":
            raise RuntimeError("summary request failed")
        return SimpleNamespace(
            sheets=(SimpleNamespace(sheet_name="Data"),),
            cache_hit=False,
        )

    pipeline._prepare_index = prepare_index
    # Make result consumption deterministic: the failed future is observed
    # before the successful future, reproducing the original progress bug.
    monkeypatch.setattr(
        query_pipeline_module,
        "as_completed",
        lambda futures: iter(futures),
    )
    samples = [
        EvalSample(
            0,
            "broken",
            "broken",
            "",
            "q",
            [],
            table_path=str(tmp_path / "broken.xlsx"),
        ),
        EvalSample(
            1,
            "working",
            "working",
            "",
            "q",
            [],
            table_path=str(tmp_path / "working.xlsx"),
        ),
    ]

    with pytest.raises(RuntimeError, match="summary request failed"):
        pipeline.prepare_samples(samples)

    assert progress_events == [
        (
            "structure_done",
            {
                "sample": "working",
                "workbook": "working.xlsx",
                "sheet": "Data",
                "status": "indexed",
            },
        )
    ]


def test_query_scoped_flow_uses_last_candidate_after_retries_exhausted(
    tmp_path: Path,
):
    workbook_path = tmp_path / "revenue.xlsx"
    _workbook(workbook_path)
    sample = _sample(workbook_path, "question-rejected-structure")
    config = _config(tmp_path)
    config["table_agent"]["query_structure_max_retries"] = 2
    llm = ChangingQueryScopedLLM()
    pipeline = TableAgentPipeline(llm, None, config)
    rejecting_verifier = AlwaysRejectVerifier()
    pipeline.verifier = rejecting_verifier

    def run_qa(**kwargs):
        structure = yaml.safe_load(
            kwargs["structure_path"].read_text(encoding="utf-8")
        )
        assert structure["revenue"]["name"] == "Annual revenue attempt 3"
        assert structure["revenue"]["headers"][0]["header_range"] == "A1"
        return (
            LLMResponse(content="200", prompt_tokens=11, completion_tokens=1),
            {"success": True, "execution_time": 0.01},
        )

    pipeline._run_verified_qa = run_qa
    pipeline.prepare_samples([sample])
    output = pipeline.run(sample)

    verification = output.metadata["verification"]["Data"]
    assert llm.structure_calls == 3
    assert rejecting_verifier.calls == 3
    assert verification["retry_exhausted"] is True
    assert verification["used_unverified_structure"] is True
    assert output.predicted_answer == "200"
