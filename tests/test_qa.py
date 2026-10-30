"""Offline tests of the QA stage: workbook tools, checks, notebook, and the answer loop with a scripted model."""

from __future__ import annotations

import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace

import openpyxl
from openpyxl.styles import Font, PatternFill
import pytest
import yaml

from TableAgent.stages.qa import orchestrator as orchestrator_module
from TableAgent.stages.qa.checker import reverse_facts, same_values
from TableAgent.stages.qa.orchestrator import answer_question
from TableAgent.stages.qa.ledger import Evidence, Need, Ledger
from TableAgent.stages.qa.prompts import (
    INSPECTOR_SYSTEM_PROMPT,
    INSPECTOR_SYSTEM_PROMPT_WITH_LOOK,
    REASONER_SYSTEM_PROMPT,
    SYNTHESIZER_SYSTEM_PROMPT,
)
from TableAgent.stages.qa.synthesizer import part_text
from TableAgent.stages.qa.tools.headers import read_layout
from TableAgent.stages.qa.tools.notebook import Notebook
from TableAgent.stages.qa.tools.sheet import SheetTools, load_structure
from TableAgent.stages.qa.utils.llm import parse_json


@pytest.fixture
def workbook_files(tmp_path: Path) -> tuple[Path, Path]:
    """Two sections (North, South) under one header row; the label "Alpha" occurs in both."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Staff by region"
    sheet.merge_cells("A1:D1")
    for col, label in zip("BCD", ("Region", "2018", "2019")):
        sheet[f"{col}3"] = label
    sheet["A4"] = "North"
    rows = {5: ("Alpha", 10, 12), 6: ("Beta", 20, 25), 8: ("Alpha", 30, 35), 9: ("Gamma", 40, 45)}
    for row, values in rows.items():
        for col, value in zip("BCD", values):
            sheet[f"{col}{row}"] = value
    sheet["A7"] = "South"
    workbook_path = tmp_path / "staff.xlsx"
    workbook.save(workbook_path)
    structure_path = tmp_path / "structure.yaml"
    structure_path.write_text(yaml.safe_dump({
        "staff": {
            "id": "staff",
            "name": "Staff by region",
            "sheet": "Data",
            "headers": [
                {"id": "region", "label": "Region", "header_range": "B3", "data_range": "B5:B9", "sub_headers": []},
                {"id": "y2018", "label": "2018", "header_range": "C3", "data_range": "C5:C9", "sub_headers": []},
                {"id": "y2019", "label": "2019", "header_range": "D3", "data_range": "D5:D9", "sub_headers": [],
                 "description": "Staff at the end of 2019"},
            ],
            "groups": [
                {"id": "north", "label": "North", "group_range": "A4", "data_range": "B5:D6"},
                {"id": "south", "label": "South", "group_range": "A7", "data_range": "B8:D9", "axis": "region"},
            ],
        }
    }), encoding="utf-8")
    return workbook_path, structure_path


@pytest.fixture
def tools(workbook_files: tuple[Path, Path]) -> SheetTools:
    workbook_path, structure_path = workbook_files
    return SheetTools(openpyxl.load_workbook(workbook_path, data_only=True), load_structure([structure_path]))


class ScriptedLLM:
    """Replies from per-role scripts; records every prompt (and each Checker system prompt) by role."""

    def __init__(
        self, reason: list[dict], inspect: list[dict] = (), synthesize: list[dict] = (), check: list[dict] = (),
        look: list[str] = (),
    ):
        self.scripts = {
            "reason": list(reason), "inspect": list(inspect), "synthesize": list(synthesize), "check": list(check),
            "look": list(look),
        }
        self.prompts: dict[str, list[str]] = {role: [] for role in self.scripts}
        self.check_systems: list[str] = []

    def generate(self, prompt: str, system_prompt: str | None = None):
        roles = {
            REASONER_SYSTEM_PROMPT: "reason", INSPECTOR_SYSTEM_PROMPT: "inspect", INSPECTOR_SYSTEM_PROMPT_WITH_LOOK: "inspect",
            SYNTHESIZER_SYSTEM_PROMPT: "synthesize",
        }
        role = "check" if system_prompt.startswith("You are a Checker") else roles[system_prompt]
        if role == "check":
            self.check_systems.append(system_prompt)
        self.prompts[role].append(prompt)
        return SimpleNamespace(content=json.dumps(self.scripts[role].pop(0)), prompt_tokens=10, completion_tokens=5)

    def generate_with_image(self, prompt: str, image_path: Path, system_prompt: str | None = None):
        assert Path(image_path).exists()
        self.prompts["look"].append(prompt)
        return SimpleNamespace(content=self.scripts["look"].pop(0), prompt_tokens=10, completion_tokens=5)


START = {
    "analysis": "One value is asked for.",
    "hypothesis": "The answer is the 2019 value of the item in its section.",
    "answer_form": "a number",
    "answer_slots": [{"slot": "the 2019 value", "asked_by": "value in 2019"}],
    "needs": [{"id": "n1", "description": "the 2019 value of Alpha in the South section"}],
}
INSPECT_N1 = {"assessment": "n1 is open.", "action": "inspect", "need_id": "n1", "strategy": "search_sheet",
              "instruction": "Alpha occurs twice; take the South one.", "reason": "need n1"}
SYNTHESIZE = {"assessment": "n1 is found.", "action": "synthesize", "reason": "done"}
ACCEPT = {"assessment": "Fits.", "verdict": "accept", "problem": "", "instruction": ""}
LOOK = {"thought": "Find the label.", "code": "print(sheet.find('Alpha'))", "done": False}
SOUTH_EVIDENCE = {
    "thought": "Take the South row.",
    "code": "evidence_n1 = {'value': 35, 'cells': ['D8'], 'where': 'South > Alpha | 2019', 'note': 'not the North row D5'}",
    "done": True,
}
PARTS = {"thought": "One slot.", "code": "answer_parts = [evidence_n1['value']]"}
REFUTED = {"verdict": "refuted", "reason": "the North row fits the need", "cells": ["D5"]}
CONFIRMED = {"verdict": "confirmed", "reason": "D8 is South's 2019 value", "cells": ["D8"]}


def _run(
    tmp_path: Path, workbook_files: tuple[Path, Path], llm: ScriptedLLM,
    question: str = "What is Alpha's value in 2019 in the South?", **settings,
):
    workbook_path, structure_path = workbook_files
    return answer_question(llm, question, workbook_path, [structure_path], tmp_path / "runs", **settings)


def _ruled(ruling: str, action: dict) -> dict:
    return {**action, "rulings": [{"need_id": "n1", "ruling": ruling, "why": f"{ruling} it"}]}


# ---------------------------------------------------------------------------------------------- workbook tools

def test_reads_resolve_merged_cells_and_find_matches_whole_words(tools: SheetTools):
    assert tools.read("A1:B1") == [["Staff by region", "Staff by region"]]
    assert tools.read("'Data'!C5") == [[10]]
    assert [ref for ref, _ in tools.find("alpha")] == ["B5", "B8"]
    assert tools.find("alp") == []
    assert tools.frame().loc[8, "D"] == 35


def test_headers_walk_to_the_row_label_column_header_section_and_structure(tools: SheetTools):
    found = tools.headers("D8")
    assert found["row_labels"] == ["Alpha"]
    assert found["column_headers"] == ["2019"] and found["headers_from"] == "structure.yaml"
    assert found["sections"] == [{"label": "South", "row": 7, "rows": "8-9", "structure": ["group: South"]}]
    assert found["siblings"] == [{"label": "North", "row": 4, "rows": "5-6", "structure": ["group: North"]}]
    assert "group: South" in found["structure"] and "header: 2019" in found["structure"]
    assert found["uncertain"] == []
    assert tools.describe("D5").startswith("D5 = 12 | row: Alpha | column: 2019 | section: North (row 4, group: North)")
    assert tools.around("D8", rows=1, cols=1).splitlines() == [
        "row 7: (empty)", "row 8: C8=30 | *D8=35", "row 9: C9=40 | D9=45",
    ]


def test_headers_list_parent_sections_but_not_siblings_and_bold_is_the_higher_level():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for ref, value in {"A2": "Region A", "B3": "Urban", "B4": "x", "C4": 5, "B5": "Rural", "B6": "y", "C6": 7}.items():
        sheet[ref] = value
    assert [section["label"] for section in SheetTools(workbook).headers("C6")["sections"]] == ["Rural", "Region A"]
    same_column = openpyxl.Workbook()
    sheet = same_column.active
    for ref, value in {"A1": "All", "A2": "Group", "A3": "x", "B3": 5}.items():
        sheet[ref] = value
    sheet["A1"].font = Font(bold=True)
    assert [section["label"] for section in SheetTools(same_column).headers("B3")["sections"]] == ["Group", "All"]


def test_a_section_label_merged_across_the_row_is_still_a_section():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"], sheet["B1"], sheet["C1"], sheet["D1"] = "Item", "Code", "Unit", "Value"
    for row, label, item in ((2, "Cafes", ("Alpha", "a", "kg", 5)), (4, "Other income", ("Beta", "b", "kg", 9))):
        sheet[f"A{row}"] = label
        sheet.merge_cells(f"A{row}:C{row}")
        sheet[f"A{row + 1}"], sheet[f"B{row + 1}"], sheet[f"C{row + 1}"], sheet[f"D{row + 1}"] = item
    tools = SheetTools(workbook)
    assert "section: Other income (row 4)" in tools.describe("D5")
    assert tools.headers("D3")["sections"][0]["rows"] == "3-3"


def test_headers_carry_merged_keys_down_and_ignore_striping(tmp_path: Path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"], sheet["B1"], sheet["C1"] = "Location", "Item", "Qty"
    for row, (location, item, qty) in enumerate(
        [("North", "Alpha", 1), (None, "Beta", 2), (None, "Gamma", 3), ("South", "Delta", 4), (None, "Eps", 5)], start=2,
    ):
        sheet[f"A{row}"], sheet[f"B{row}"], sheet[f"C{row}"] = location, item, qty
        sheet[f"B{row}"].fill = PatternFill("solid", fgColor="FFDDEEFF" if row % 2 else "FFFFFFFF")
    sheet.merge_cells("A2:A4")
    sheet["B7"], sheet["C7"] = "Total", "=SUM(C2:C6)"
    path = tmp_path / "keys.xlsx"
    workbook.save(path)
    tools = SheetTools(openpyxl.load_workbook(path, data_only=True), formula_path=path)

    found = tools.headers("C3")
    assert found["block"] == {"rows": "2-6", "starts": "row 1 is the header", "ends": "row 7 has no data (label 'Total')"}
    assert {"header": "Location", "value": "North", "from": "A2", "rows": "2-4", "next": "South"}.items() <= found["keys"][0].items()
    assert "key: Location = North (from A2)" in tools.describe("C3")
    assert read_layout(openpyxl.load_workbook(path).active, 1, openpyxl.load_workbook(path).active).totals == {7: (2, 6)}


def test_structure_find_matches_labels_first_then_descriptions(tools: SheetTools):
    assert tools.structure_find("south") == [{
        "kind": "group", "path": "South", "sheet": "Data", "label_range": "A7", "data_range": "B8:D9",
        "description": "", "axis": "region", "matched": "label",
    }]
    assert [(item["path"], item["matched"]) for item in tools.structure_find("end of")] == [("2019", "description")]
    assert tools.structure_find("sou") == [] and tools.structure_find("south", sheet="Other") == []


def test_striping_in_part_of_a_sheet_does_not_split_blocks_but_a_bold_total_does():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"], sheet["B1"] = "Item", "Value"
    for row in range(2, 222):  # 20 rows striped for reading, then 200 plain rows
        sheet[f"A{row}"], sheet[f"B{row}"] = f"item {row}", row
        if row < 22 and row % 2:
            sheet[f"A{row}"].fill = PatternFill("solid", fgColor="FFDDEEFF")
    for row in (222, 223):
        sheet[f"A{row}"], sheet[f"B{row}"] = "Total" if row == 222 else "after", 1
    sheet["A222"].font = Font(bold=True)
    tools = SheetTools(workbook)

    assert tools.headers("B6")["block"]["rows"] == "2-221"
    assert tools.headers("B6")["block"]["ends"] == "row 222 is formatted differently (bold)"


def test_an_unsure_walk_asks_the_vision_model_about_the_area_a_few_times():
    workbook = openpyxl.Workbook()
    for row in range(1, 6):
        workbook.active[f"A{row}"] = row
    asked = []
    tools = SheetTools(workbook, look=lambda ref, question, sheet: asked.append(ref) or "a title row above")

    found = tools.headers("A5")
    assert found["uncertain"][0] == "no column header found above the cell"
    assert found["vision"] == "a title row above" and asked == ["A1:A5"]
    assert "vision" not in tools.headers("A4", look=False)
    for _ in range(3):
        tools.headers("A3")
    assert len(asked) == 3  # at most three images per question


def test_formats_list_the_format_and_formula_of_filled_cells(tmp_path: Path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"], sheet["B1"], sheet["A2"], sheet["B2"], sheet["A3"], sheet["B3"] = "Item", None, "x", 4, "Total", "=SUM(B2:B2)"
    sheet["A3"].font = Font(bold=True)
    path = tmp_path / "f.xlsx"
    workbook.save(path)
    tools = SheetTools(openpyxl.load_workbook(path), formula_path=path)
    assert tools.formats("A1:B3").splitlines() == [
        "A1 = Item | plain", "A2 = x | plain", "B2 = 4 | plain", "A3 = Total | bold", "B3 = =SUM(B2:B2) | plain | =SUM(B2:B2)",
    ]


def test_look_is_unavailable_without_a_vision_model(tools: SheetTools):
    with pytest.raises(ValueError, match="not available"):
        tools.look("B4:D8", "Which cells are shaded?")
    assert "vision" not in tools.headers("A2")


# ------------------------------------------------------------------------------------------------------ checks

def test_reverse_facts_give_the_meaning_of_the_cells_and_their_values(tools: SheetTools):
    ledger = Ledger("q")
    facts = reverse_facts(tools, ledger, Evidence("n1", 1, 1, "search_sheet", "", True, 35, ["D8"]))
    assert facts[1].strip().startswith("D8 = 35 | row: Alpha | column: 2019 | section: South (row 7, group: South)")
    assert "the cells hold the reported values" in facts
    assert "  header 2019: Staff at the end of 2019" in facts

    facts = reverse_facts(tools, ledger, Evidence("n1", 1, 1, "search_sheet", "", True, [12, 99], ["D5", "D8"]))
    assert "values differ: D8 holds '35', reported '99'" in facts
    assert "2 filled cell(s) in the column within rows 5-6, 8-9 were not taken:" in facts
    assert any(line.strip().startswith("D6 = 25 | row: Beta") for line in facts)
    assert "4 value(s) reported but 2 cell(s) given" in " ".join(
        reverse_facts(tools, ledger, Evidence("n1", 1, 1, "s", "", True, [1, 2, 3, 4], ["D5", "D8"]))
    )


def test_reverse_facts_name_cells_shared_with_another_need(tools: SheetTools):
    ledger = Ledger("q")
    ledger.start({**START, "needs": [{"id": "n1", "description": "a"}, {"id": "n2", "description": "b"}]})
    ledger.record(Evidence("n1", 1, 1, "search_sheet", "", True, [12, 25], ["D5:D6"]))
    facts = reverse_facts(tools, ledger, Evidence("n2", 1, 2, "search_sheet", "", True, 25, ["D6"]))
    assert "shared with the evidence of n1: D6 (combining both counts them twice)" in facts


def test_values_are_the_same_in_any_order_and_form():
    assert same_values({"Alpha": 12, "Beta": "25"}, [25.0, 12])
    assert not same_values([12, 25], [12])
    assert not same_values("North", "South")


# ---------------------------------------------------------------------------------------------------- notebook

def test_notebook_keeps_variables_and_refuses_system_modules():
    notebook = Notebook({})
    assert notebook.run("x = 2\nprint(x * 3)", "a").output == "6"
    assert notebook.run("print(x + 1)", "b").output == "3"
    assert "not allowed" in notebook.run("import os", "c").error
    failed = notebook.run("y = 1\nz = y / 0", "d")
    assert failed.error.splitlines() == ["line 2: z = y / 0", "ZeroDivisionError: division by zero"]


def test_notebooks_running_in_threads_keep_their_own_output():
    notebooks = [Notebook({}) for _ in range(8)]
    code = "import time\nfor i in range(5):\n    print(name, end='')\n    time.sleep(0.02)"
    for index, notebook in enumerate(notebooks):
        notebook.namespace["name"] = str(index)
    with ThreadPoolExecutor(8) as pool:
        cells = list(pool.map(lambda notebook: notebook.run(code, "a"), notebooks))
    assert [cell.output for cell in cells] == [str(index) * 5 for index in range(8)]


def test_notebook_allows_strptime_but_not_its_module():
    notebook = Notebook({})
    assert notebook.run("import datetime\nprint(datetime.datetime.strptime('2020-01-31', '%Y-%m-%d').day)", "a").output == "31"
    assert "not allowed" in notebook.run("import _strptime", "b").error


# ------------------------------------------------------------------------------------------------- answer loop

def test_answers_through_reasoning_inspection_synthesis_and_end(tmp_path, workbook_files):
    llm = ScriptedLLM([START, INSPECT_N1, SYNTHESIZE, ACCEPT], [LOOK, SOUTH_EVIDENCE], [PARTS])
    result = _run(tmp_path, workbook_files, llm)

    assert result.success and result.answer == "35"
    assert "[('B5', 'Alpha'), ('B8', 'Alpha')]" in llm.prompts["inspect"][1]  # turn 2 sees what turn 1 printed
    assert "where: South > Alpha | 2019" in llm.prompts["reason"][2]
    assert result.ledger.evidence[0].checks == []  # no checks run by default
    trace = Path(result.artifacts["trace_md"]).read_text(encoding="utf-8")
    assert "#### Inspection n1, attempt 1 (search_sheet): found" in trace
    assert "### Synthesis → 35" in trace and "### Abstract reasoning: end (accept)" in trace
    ledger = json.loads(Path(result.artifacts["ledger_json"]).read_text(encoding="utf-8"))
    assert ledger["needs"]["n1"]["status"] == "found"
    assert sorted(path.name for path in Path(result.artifacts["cells_dir"]).iterdir())[0] == "01_n1_attempt1_turn1.py"
    events = [json.loads(line)["event_type"] for line in Path(result.artifacts["events_jsonl"]).read_text().splitlines()]
    assert events[0] == "question" and events[-1] == "result"
    readable = Path(result.artifacts["events_log"]).read_text(encoding="utf-8")
    assert "llm_call #1 ok" in readable and "prompt: (" in readable
    assert "inspection_turn\n  need_id: n1\n  attempt: 1\n  number: 1\n  thought: Find the label." in readable


def test_synthesis_is_refused_while_a_need_is_open(tmp_path, workbook_files):
    llm = ScriptedLLM([START, SYNTHESIZE, INSPECT_N1, SYNTHESIZE, ACCEPT], [SOUTH_EVIDENCE], [PARTS])
    result = _run(tmp_path, workbook_files, llm)

    assert result.success
    assert "Cannot synthesize while needs are open: n1" in llm.prompts["reason"][2]
    assert result.ledger.decisions[0].accepted is False


def test_an_error_turn_is_an_observation_and_done_needs_the_evidence(tmp_path, workbook_files):
    broken = {"thought": "Try.", "code": "print(undefined_name)", "done": False}
    early = {"thought": "Done?", "code": "print('ready')", "done": True}
    llm = ScriptedLLM([START, INSPECT_N1, SYNTHESIZE, ACCEPT], [broken, early, SOUTH_EVIDENCE], [PARTS])
    result = _run(tmp_path, workbook_files, llm)

    assert "NameError: name 'undefined_name' is not defined" in llm.prompts["inspect"][1]
    assert "done was refused: `evidence_n1` is not set to a dict." in llm.prompts["inspect"][2]
    assert result.answer == "35"


def test_the_last_inspection_turn_asks_for_the_evidence(tmp_path, workbook_files):
    llm = ScriptedLLM([START, INSPECT_N1, SYNTHESIZE, ACCEPT], [LOOK] * 4 + [SOUTH_EVIDENCE], [PARTS])
    result = _run(tmp_path, workbook_files, llm)

    assert "This is the last turn" in llm.prompts["inspect"][4] and "last turn" not in llm.prompts["inspect"][3]
    assert result.answer == "35"


def test_synthesis_retries_until_there_is_one_part_per_slot(tmp_path, workbook_files):
    two = {"thought": "Two.", "code": "answer_parts = [35, 'extra']"}
    llm = ScriptedLLM([START, INSPECT_N1, SYNTHESIZE, ACCEPT], [SOUTH_EVIDENCE], [two, PARTS])
    result = _run(tmp_path, workbook_files, llm)

    assert "has 2 entries; it needs exactly 1" in llm.prompts["synthesize"][1]
    assert result.answer == "35"


def test_answer_parts_are_written_plainly():
    assert part_text(True) == "Yes" and part_text(16.0) == "16" and part_text(2.5) == "2.5"
    assert part_text(["A", 3.0]) == "A, 3"
    assert part_text(10.15 + 10.15) == "20.3" and part_text(0.30000000000000004) == "0.3"


def test_a_returned_review_reopens_exploration_once(tmp_path, workbook_files):
    returned = {"assessment": "Wrong section.", "verdict": "return", "problem": "North value", "instruction": "use South"}
    revise = {"assessment": "Revise.", "action": "revise_need", "need_id": "n1",
              "description": "the 2019 value of Alpha under the South section", "reason": "review"}
    llm = ScriptedLLM(
        [START, INSPECT_N1, SYNTHESIZE, returned, revise, INSPECT_N1, SYNTHESIZE],
        [SOUTH_EVIDENCE, SOUTH_EVIDENCE], [PARTS, PARTS],
    )
    result = _run(tmp_path, workbook_files, llm)

    assert result.success and result.ledger.returns == 1
    assert "The review returned the answer '35'" in llm.prompts["reason"][4]
    assert result.ledger.needs["n1"].inspections_left == 2  # revising restarts the inspection count
    assert len(result.ledger.answers) == 2  # the second answer is final without another review


def test_running_out_of_steps_with_an_open_need_fails(tmp_path, workbook_files, monkeypatch):
    monkeypatch.setattr(orchestrator_module, "MAX_STEPS", 2)
    llm = ScriptedLLM([START, SYNTHESIZE, SYNTHESIZE])
    result = _run(tmp_path, workbook_files, llm)

    assert not result.success and result.error == "Out of budget with open needs: n1."
    assert Path(result.artifacts["trace_md"]).is_file()


def test_look_renders_the_range_for_the_vision_model(tmp_path, workbook_files):
    class Renderer:
        def render_workbook_range(self, workbook_path, sheet, area, image_path):
            image_path.parent.mkdir(parents=True, exist_ok=True)
            image_path.write_bytes(b"png")

    look = {"thought": "Look at the marks.", "code": "print(sheet.look('B4:D8', 'Which cells are shaded?'))", "done": False}
    llm = ScriptedLLM([START, INSPECT_N1, SYNTHESIZE, ACCEPT], [look, SOUTH_EVIDENCE], [PARTS], look=["D8 is shaded."])
    result = _run(tmp_path, workbook_files, llm, vision_client=llm, renderer=Renderer())

    assert result.success
    assert 'range B4:D8 of sheet "Data"' in llm.prompts["look"][0]
    assert "D8 is shaded." in llm.prompts["inspect"][1]


# ------------------------------------------------------------------------------------------------------ ledger

def test_ledger_rules():
    ledger = Ledger("q")
    ledger.start({**START, "needs": [{"id": "n1", "description": "a"}, {"id": "n2", "description": "b"}]})
    assert ledger.refusal("inspect", {"need_id": "n1", "strategy": "guess"}).startswith("Unknown strategy")
    assert ledger.refusal("drop_need", {"need_id": "n1"}) == ""
    ledger.apply("drop_need", {"need_id": "n1"}, 1, "not needed")
    assert ledger.refusal("drop_need", {"need_id": "n2"}).startswith("Cannot drop n2, the last need")
    for attempt in (1, 2, 3):
        ledger.record(Evidence("n2", attempt, attempt, "search_sheet", ""))
    assert "inspected 3 times" in ledger.refusal("inspect", {"need_id": "n2", "strategy": "window"})
    ledger.apply("revise_need", {"need_id": "n2", "description": "b, in rows 2-9"}, 4, "narrow it")
    assert "inspect it before revising it again" in ledger.refusal("revise_need", {"need_id": "n2", "description": "c"})
    assert ledger.refusal("synthesize", {}).startswith("Cannot synthesize while needs are open: n2")
    assert ledger.refusal("revise_hypothesis", {"hypothesis": "h", "answer_slots": [{"slot": "x"}]}).startswith(
        "each answer slot"
    )
    assert ledger.refusal("jump", {}).startswith("Unknown action 'jump'")
    assert Need("n9", "x").variable == "evidence_n9"


def test_parse_json_takes_code_from_python_block_and_raw_line_breaks():
    assert parse_json('{"thought": "two\nlines", "done": false}\n```python\nprint(sheet.find("Alpha"))\n```') == {
        "thought": "two\nlines", "done": False, "code": 'print(sheet.find("Alpha"))',
    }


def test_parse_json_repairs_unescaped_quotes_inside_strings():
    reply = '{"thought": "the "Alpha" row, then sheet.read("B3:D9")", "done": false}'
    assert parse_json(reply) == {"thought": 'the "Alpha" row, then sheet.read("B3:D9")', "done": False}


# ----------------------------------------------------------------------------------------------------- checkers

def test_a_refuting_check_must_be_ruled_on_before_synthesis_then_shows_in_one_line(tmp_path, workbook_files):
    llm = ScriptedLLM(
        [START, INSPECT_N1, SYNTHESIZE, _ruled("keep", SYNTHESIZE), ACCEPT], [SOUTH_EVIDENCE], [PARTS], [REFUTED],
    )
    result = _run(tmp_path, workbook_files, llm, checks=("reverse",))

    assert "D8 = 35 | row: Alpha | column: 2019 | section: South" in llm.prompts["check"][0]
    step = llm.prompts["reason"][2]
    assert "check reverse: REFUTED - the North row fits the need (cells: D5)" in step and "RULE ON IT" in step
    assert "Checks to rule on: n1." in step
    assert "Cannot synthesize before ruling on the checks of: n1." in llm.prompts["reason"][3]
    assert "check reverse: refuted - the North row fits the need; ruled keep: keep it" in llm.prompts["reason"][4]
    assert result.success and result.answer == "35"
    trace = Path(result.artifacts["trace_md"]).read_text(encoding="utf-8")
    assert "##### Check reverse: refuted" in trace and "**Ruling on n1:** keep" in trace


def test_rejecting_the_evidence_reopens_the_need(tmp_path, workbook_files):
    reinspect = _ruled("reject", INSPECT_N1)
    llm = ScriptedLLM(
        [START, INSPECT_N1, reinspect, SYNTHESIZE, ACCEPT], [SOUTH_EVIDENCE, SOUTH_EVIDENCE], [PARTS], [REFUTED, CONFIRMED],
    )
    result = _run(tmp_path, workbook_files, llm, checks=("reverse",))

    assert result.ledger.evidence[0].ruling["ruling"] == "reject" and len(result.ledger.evidence) == 2
    assert "rejected: reject it" in llm.prompts["reason"][3]
    assert "check reverse: confirmed - D8 is South's 2019 value" in llm.prompts["reason"][3]
    assert result.answer == "35"


def test_a_later_check_runs_only_when_the_earlier_did_not_confirm(tmp_path, workbook_files):
    llm = ScriptedLLM([START, INSPECT_N1, SYNTHESIZE, ACCEPT], [SOUTH_EVIDENCE], [PARTS], [CONFIRMED])
    result = _run(tmp_path, workbook_files, llm, checks=("reverse", "debate"))

    assert [check["by"] for check in result.ledger.evidence[0].checks] == ["reverse"]


def test_debate_objections_must_name_cells(tmp_path, workbook_files):
    probe = {"thought": "Other Alpha rows.", "code": "print(sheet.describe('D5'))", "done": False}
    objections = {"thought": "Two doubts.", "done": True, "objections": [
        {"objection": "the North Alpha row also fits", "cells": ["D5"]}, {"objection": "a feeling", "cells": []},
    ]}
    llm = ScriptedLLM(
        [START, INSPECT_N1, _ruled("keep", SYNTHESIZE), ACCEPT], [SOUTH_EVIDENCE], [PARTS], [probe, objections],
    )
    result = _run(tmp_path, workbook_files, llm, checks=("debate",))

    step = llm.prompts["reason"][2]
    assert "check debate: REFUTED - 1 objection(s) (cells: D5)" in step
    assert "- the North Alpha row also fits (D5 = 12)" in step and "a feeling" not in step
    assert "D5 = 12 | row: Alpha" in llm.prompts["check"][1]  # the debate has its own notebook with `claim`
    assert "n1_attempt1_debate_turn1.py" in Path(result.artifacts["trace_md"]).read_text(encoding="utf-8")


def test_a_blind_independent_check_reports_its_own_value_and_the_runner_compares(tmp_path, workbook_files):
    own = {"thought": "North row.", "done": True, "code": "found = {'value': 12, 'cells': ['D5']}", "reason": "row 5"}
    llm = ScriptedLLM([START, INSPECT_N1, _ruled("keep", SYNTHESIZE), ACCEPT], [SOUTH_EVIDENCE], [PARTS], [own])
    result = _run(tmp_path, workbook_files, llm, checks=("independent",), sees_value=False)

    check = result.ledger.evidence[0].checks[0]
    assert check["verdict"] == "refuted" and check["reason"] == "the Checker read 12 from D5 (row 5)"
    assert "`claim`" not in llm.check_systems[0] and "structure.yaml is not available" in llm.check_systems[0]
    assert "Inspector's evidence" not in llm.prompts["check"][0]
    assert "Workbook preview" in llm.prompts["check"][0] and "Alpha" in llm.prompts["check"][0]


def test_an_independent_check_sees_only_its_view_and_its_last_turn_gives_the_verdict(tmp_path, workbook_files):
    probe = {"thought": "Which tools?", "code": "print(hasattr(sheet, 'headers'), claim['value'], sheet.formats('A7:D8'))", "done": False}
    verdict_with_code = {**CONFIRMED, "thought": "It holds.", "done": False, "code": "print('not run')"}
    llm = ScriptedLLM(
        [START, INSPECT_N1, SYNTHESIZE, ACCEPT], [SOUTH_EVIDENCE], [PARTS], [probe] * 3 + [verdict_with_code],
    )
    result = _run(tmp_path, workbook_files, llm, checks=("independent",))

    assert "False 35 A7 = South | plain\nB8 = Alpha | plain" in llm.prompts["check"][1]
    assert "This is the last turn" in llm.prompts["check"][3]
    assert result.ledger.evidence[0].checks[0]["verdict"] == "confirmed"
    assert "check independent: confirmed - D8 is South's 2019 value" in llm.prompts["reason"][2]


def test_the_method_check_reaches_the_final_review(tmp_path, workbook_files):
    flawed = {"verdict": "flawed", "problem": "the question asks for a list, the code sums"}
    llm = ScriptedLLM([START, INSPECT_N1, SYNTHESIZE, ACCEPT], [SOUTH_EVIDENCE], [PARTS], [flawed])
    result = _run(tmp_path, workbook_files, llm, synthesis_check=True)

    assert "answer_parts = [evidence_n1['value']]" in llm.prompts["check"][0]
    assert "found it flawed: the question asks for a list, the code sums" in llm.prompts["reason"][-1]
    assert "**Method check:** flawed" in Path(result.artifacts["trace_md"]).read_text(encoding="utf-8")
